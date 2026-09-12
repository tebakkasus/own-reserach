#!/usr/bin/env python3
"""
Build Full-Text Indonesian Manuscript (v8 Final) for Study & Learning.
Translates all updated v8 content (HKSJ statistics, associational wording, SCM objective framing,
search limitations, 14 vs 3 study clarification, A4 layout, declarations) into clear academic Bahasa Indonesia.
"""
import docx
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

OUTPUT = "Manuscript_IKH_2026_BAHASA_INDONESIA_v8.docx"

doc = docx.Document()
for section in doc.sections:
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def add_p(text="", bold=False, italic=False, size=11, align=None, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p

# ============================================================
# HALAMAN JUDUL
# ============================================================
add_p("Artikel Penelitian", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

add_p("Efektivitas Intervensi Komunitas dalam Menurunkan Dispensing Antibiotik Tanpa Resep di Apotek Negara Berkembang: Meta-Analisis",
      bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.15)

p_auth = add_p("", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
r1 = p_auth.add_run("Tengku Muhammad Lufthi Hannur")
r1.bold = True; r1.font.name = 'Times New Roman'; r1.font.size = Pt(11)
r2 = p_auth.add_run("1*")
r2.bold = True; r2.font.name = 'Times New Roman'; r2.font.size = Pt(10); r2.font.superscript = True

p_aff = add_p("", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
ra1 = p_aff.add_run("1 ")
ra1.font.name = 'Times New Roman'; ra1.font.size = Pt(8.5); ra1.font.superscript = True
ra2 = p_aff.add_run("Program Studi Pendidikan Dokter, Fakultas Kedokteran, Universitas Islam Sumatera Utara, Medan, Sumatera Utara, Indonesia")
ra2.font.name = 'Times New Roman'; ra2.font.size = Pt(9.5)

add_p("* Penulis Korespondensi: Tengku Muhammad Lufthi Hannur | Email: tengkumuhammadlufthihannur@fk.uisu.ac.id | WhatsApp: +62 822-6740-3241 | ORCID: https://orcid.org/0009-0008-1306-2485",
      italic=True, size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

doc.add_page_break()

# ============================================================
# ABSTRAK BAHASA INDONESIA
# ============================================================
add_p("ABSTRAK", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

abstract_sections = [
    ("Latar Belakang: ", "Dispensing antibiotik tanpa resep (NPD) di apotek komunitas berkontribusi terhadap resistensi antimikroba, namun belum ada tinjauan kuantitatif yang menggabungkan estimasi efek intervensi khusus dari apotek komunitas swasta di negara berkembang (LMIC). "),
    ("Metode: ", "Tinjauan sistematis dan meta-analisis ini mengikuti panduan PRISMA 2020. Penelusuran dilakukan pada PubMed, OpenAlex, dan Cochrane Central Register of Controlled Trials hingga September 2026. Studi yang memenuhi syarat adalah studi intervensi terkontrol yang mengevaluasi intervensi untuk menurunkan NPD di apotek komunitas swasta di LMIC, dengan luaran diukur melalui metode pasien tersimulasi tanpa pemberitahuan (unannounced simulated client methods). Odds ratio ter-adjust klaster digabungkan menggunakan model efek acak (random-effects model) dengan estimator DerSimonian-Laird dan pembobotan inverse-variance. Analisis sensitivitas Hartung-Knapp-Sidik-Jonkman dilakukan untuk menilai ketahanan interval kepercayaan mengingat terbatasnya jumlah studi. "),
    ("Hasil: ", "Tiga studi terkontrol dari Vietnam, Nigeria, dan Indonesia memenuhi kriteria inklusi sintesis kuantitatif (k = 3; 345 apotek; 1.683 kunjungan pasien tersimulasi). Peluang NPD pada kelompok intervensi 83,6% lebih rendah dibandingkan kelompok kontrol (pooled OR = 0,164; IK 95%: 0,102 hingga 0,265; Z = -7,38; p < 0,0001). Metode Hartung-Knapp-Sidik-Jonkman menghasilkan interval yang lebih lebar (OR = 0,164; IK 95%: 0,064 hingga 0,419; t(2) = -8,31, p = 0,014). Heterogenitas statistik tidak terdeteksi (I-squared = 0,0%; Q = 1,58, df = 2, p = 0,454). "),
    ("Kesimpulan: ", "Sepengetahuan kami, ini merupakan meta-analisis pertama yang menggabungkan data studi intervensi terkontrol mengenai penurunan NPD di apotek komunitas swasta LMIC. Intervensi multifaset yang menggabungkan edukasi petugas apotek, uji diagnostik cepat, dan pengawasan sejawat berhubungan dengan peluang dispensing antibiotik tanpa resep yang lebih rendah.")
]

for label, text in abstract_sections:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    rl = p.add_run(label); rl.bold = True; rl.font.name='Times New Roman'; rl.font.size = Pt(10.5)
    rt = p.add_run(text); rt.font.name='Times New Roman'; rt.font.size = Pt(10.5)

p_kw = doc.add_paragraph()
p_kw.paragraph_format.space_after = Pt(14)
rkl = p_kw.add_run("Kata Kunci: "); rkl.bold = True; rkl.font.name='Times New Roman'; rkl.font.size = Pt(9.5)
rkv = p_kw.add_run("Resistensi Antimikroba; Apotek Komunitas; Dispensing Antibiotik Tanpa Resep; Meta-Analisis; Negara Berkembang")
rkv.font.name='Times New Roman'; rkv.font.size = Pt(9.5)

doc.add_page_break()

# ============================================================
# PENDAHULUAN
# ============================================================
add_p("Pendahuluan", bold=True, size=12, space_before=6, space_after=6)

add_p("Resistensi antimikroba (AMR) berhubungan dengan estimasi 4,95 juta kematian secara global pada tahun 2019, dengan beban terbesar dialami oleh negara berpenghasilan rendah dan menengah (LMIC). Salah satu kontributor yang dapat dimodifikasi adalah dispensing antibiotik tanpa resep (NPD), yang didefinisikan sebagai penjualan antibiotik sistemik beresep oleh sarana ritel obat swasta tanpa resep medis yang sah. Di banyak LMIC, apotek komunitas swasta merupakan titik kontak medis pertama bagi masyarakat karena lokasinya yang dekat dengan tempat tinggal, tidak memerlukan janji temu, dan tidak mengenakan biaya konsultasi seperti pada fasilitas kesehatan formal.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Tinjauan sistematis terdahulu telah mengevaluasi intervensi untuk memperbaiki dispensing antibiotik di negara berkembang, namun menghadapi keterbatasan metodologis yang penting. Tinjauan oleh Afari-Asiedu et al. (2022) hanya menyajikan sintesis naratif karena heterogenitas antar-studi menghalangi penggabungan kuantitatif. Demikian pula, tinjauan Cochrane (2025) tidak dapat menggabungkan estimasi efek akibat perbedaan desain dan setting. Kedua tinjauan tersebut menggabungkan fasilitas sektor publik (seperti puskesmas dan klinik pemerintah) dengan apotek ritel swasta. Padahal, apotek swasta dan fasilitas publik beroperasi di bawah tekanan finansial dan regulasi yang berbeda. Apotek swasta sangat bergantung pada penjualan langsung kepada konsumen, menghadapi persaingan bisnis, dan mengalami pengawasan regulasi yang relatif lebih longgar, sehingga perilaku dispensing-nya tidak dapat disamakan dengan fasilitas publik.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Oleh karena itu, kami melakukan tinjauan sistematis dan meta-analisis yang dibatasi khusus pada studi intervensi terkontrol di apotek komunitas swasta LMIC. Dengan berfokus pada satu sektor spesifik dan pada luaran yang diukur melalui metode pasien tersimulasi tanpa pemberitahuan (unannounced simulated client methods), penelitian ini bertujuan untuk memperoleh estimasi efek gabungan ter-adjust pengelompokan yang defensibel guna mendukung kebijakan tata kelola antibiotik di tingkat komunitas.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ============================================================
# METODE
# ============================================================
add_p("Metode", bold=True, size=12, space_before=6, space_after=6)

add_p("Desain Studi dan Protokol: Tinjauan sistematis dan meta-analisis ini disusun mengikuti panduan Preferred Reporting Items for Systematic Reviews and Meta-Analyses (PRISMA 2020).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Kriteria Inklusi: Studi dinyatakan memenuhi syarat apabila: (1) dilakukan di apotek komunitas swasta atau toko obat ritel swasta di negara berpenghasilan rendah dan menengah (LMIC) versi Bank Dunia; (2) mengevaluasi intervensi edukasional, manajerial, diagnostik, atau regulasi yang bertujuan menurunkan dispensing antibiotik; (3) memiliki kelompok pembanding konkuren (concurrent control group), yaitu uji acak berkelompok (cluster-randomised trial) atau studi pra-pasca terkontrol (controlled before-after study); (4) mengukur angka dispensing antibiotik tanpa resep secara kuantitatif menggunakan metode pasien tersimulasi tanpa pemberitahuan (unannounced simulated client methods); serta (5) melaporkan estimasi efek yang telah disesuaikan terhadap efek pengelompokan (cluster-adjusted effect size) atau menyediakan data dasar untuk penghitungannya.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Sumber Informasi dan Strategi Pencarian: Penelusuran sistematis dilakukan pada PubMed, OpenAlex, dan Cochrane Central Register of Controlled Trials (CENTRAL) hingga September 2026. Pencarian menggabungkan istilah kosa kata terkontrol dan kata kunci bebas. Penelusuran pelacakan sitasi (citation tracking) maju dan mundur juga dilakukan terhadap daftar pustaka studi terpilih dan tinjauan sistematis terdahulu. Pencarian dibatasi pada ketiga basis data tersebut; studi regional yang tidak terindeks mungkin tidak terekam.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Seleksi Studi dan Ekstraksi Data: Penapisan rekaman dilakukan berdasarkan judul dan abstrak, dilanjutkan dengan penilaian teks lengkap terhadap kriteria inklusi. Sintesis tinjauan sistematis mempertahankan kolam bukti kualitatif yang lebih luas (n = 14) untuk mengevaluasi tata kelola antimikroba dan praktik dispensing di setting komunitas dan ritel, sedangkan sintesis kuantitatif primer dibatasi secara ketat pada studi intervensi terkontrol di apotek komunitas swasta LMIC yang melaporkan estimasi efek ter-adjust klaster menggunakan metode pasien tersimulasi (n = 3). Data diekstraksi ke dalam matriks tersertifikasi yang mencakup nama penulis pertama, tahun, negara, desain studi, jumlah apotek, jumlah kunjungan tersimulasi, komponen intervensi, pembanding, jumlah kejadian dispensing, estimasi efek (aOR), IK 95%, nilai p, serta model statistik klaster yang digunakan (GLIMMIX, GLMER, xtlogit).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Penilaian Risiko Bias: Risiko bias dinilai menggunakan instrumen Cochrane RoB 2 khusus cluster-RCT dan instrumen ROBINS-I untuk studi non-acak (controlled before-after).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Sintesis Statistik: Besaran efek digabungkan dengan model efek acak (DerSimonian-Laird random-effects model) menggunakan pembobotan inverse-variance di R (paket metafor dan meta). Untuk setiap studi, odds ratio ter-adjust klaster (aOR) dan IK 95% ditransformasikan ke bentuk logaritma natural (yi = ln[aOR]) dan standar error (sei = {ln[IK atas] - ln[IK bawah]} / 3.92). Heterogenitas dinilai menggunakan uji Cochran's Q (p < 0.10), statistik I², dan varians antar-studi τ². Mengingat hanya tiga studi yang masuk ke sintesis kuantitatif, analisis sensitivitas Hartung-Knapp-Sidik-Jonkman (HKSJ) dilakukan untuk menilai ketahanan interval kepercayaan. Visualisasi forest plot dibuat menggunakan Matplotlib di Python 3.12.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ============================================================
# HASIL
# ============================================================
add_p("Hasil", bold=True, size=12, space_before=6, space_after=6)

add_p("Hasil Seleksi Studi: Pencarian basis data menghasilkan 3.223 rekaman, dan pelacakan sitasi menambahkan 22 rekaman, dengan total 3.245 rekaman. Setelah mengeluarkan 2.635 rekaman sebelum penapisan judul/abstrak dan melakukan evaluasi judul/abstrak (n = 610; 427 dikeluarkan), sebanyak 183 teks lengkap dinilai kelayakannya, di mana 169 teks lengkap dikeluarkan karena alasan metodologis (setting dokter sektor publik, n=18; toko obat non-swasta, n=34; desain tanpa kontrol/observasional, n=42; luaran bukan NPD, n=26; intervensi berorientasi dokter primer, n=14; manajemen malaria saja, n=9; alasan lain, n=26). Terdapat 14 studi yang memenuhi kriteria kualitatif tinjauan sistematis (n = 14), dan tiga studi di antaranya memenuhi kriteria ketat sintesis kuantitatif meta-analisis (Gambar 1).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

if os.path.exists("Plots/prisma_flow_diagram.png"):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(6)
    p_img1.paragraph_format.space_after = Pt(4)
    run1 = p_img1.add_run()
    run1.add_picture("Plots/prisma_flow_diagram.png", width=Inches(5.5))
    add_p("Gambar 1. Diagram Alur PRISMA 2020 Seleksi Studi Meta-Analisis",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

add_p("Karakteristik Studi Terintegrasi: Tiga studi terkontrol berasal dari tiga negara (Tabel 1). Chalker et al. (2005) melakukan cluster-RCT berpasangan (matched-pair) di Hanoi, Vietnam, menguji strategi 3 fase (penegakan regulasi, inspeksi, seminar edukasi, dan pengawasan sejawat) di 55 apotek swasta (273 kunjungan tersimulasi). Onwunduba et al. (2023) melakukan cluster-RCT tersarang (stratified) di Awka, Nigeria, menguji penggunaan alat uji cepat CRP (point-of-care test) dan pelatihan staf di 20 apotek swasta (600 kunjungan). Ferdiana et al. (2024) melakukan evaluasi pra-pasca terkontrol terhadap paket PINTAR (modul edukasi online, kampanye publik, edukasi sejawat, dan sertifikasi branding) di Semarang, Indonesia, pada 270 apotek swasta (80 intervensi dan 190 kontrol; 810 kunjungan). Ketiga studi menggunakan metode pasien tersimulasi tanpa pemberitahuan dengan skenario klinis terstandar (ISPA, ISK, dan diare anak).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

add_p("Tabel 1. Karakteristik Studi Terkontrol yang Diikutsertakan dalam Sintesis Kuantitatif",
      bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

table = doc.add_table(rows=4, cols=6)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Studi (Tahun)", "Negara", "Desain Studi", "Sampel (Apotek / Kunjungan)", "Komponen Intervensi", "Adjusted OR (IK 95%)"]
hdr_row = table.rows[0]
for j, h in enumerate(headers):
    cell = hdr_row.cells[j]
    cell.text = h
    cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

table_data = [
    ["Chalker et al. (2005)", "Vietnam (Hanoi)", "Cluster-randomised (matched-pair)", "55 / 273", "Penegakan regulasi, inspeksi, edukasi, peer review", "0,134 (0,057–0,316)"],
    ["Onwunduba et al. (2023)", "Nigeria (Awka)", "Cluster-randomised (stratified)", "20 / 600", "Tes cepat CRP point-of-care plus pelatihan staf", "0,279 (0,107–0,726)"],
    ["Ferdiana et al. (2024)", "Indonesia (Semarang)", "Controlled before-after (pra-pasca)", "270 / 810", "Paket PINTAR: modul online, kampanye, peer review", "0,140 (0,070–0,300)"]
]
for i, row in enumerate(table_data):
    r = table.rows[i+1]
    for j, val in enumerate(row):
        c = r.cells[j]
        c.text = val
        p = c.paragraphs[0]
        if len(p.runs) > 0:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 5] else WD_ALIGN_PARAGRAPH.LEFT

add_p("", space_after=6)

add_p("Meta-Analisis: Dalam meta-analisis model efek acak (Gambar 2), pooled odds ratio untuk NPD adalah 0,164 (IK 95%: 0,102 hingga 0,265; Z = -7,38; p < 0,0001). Hasil ini setara dengan 83,6% peluang lebih rendah untuk melakukan dispensing tanpa resep di apotek intervensi dibandingkan apotek kontrol. Estimasi aOR individual berkisar antara 0,134 hingga 0,279 dengan arah efek yang konsisten.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Analisis Sensitivitas: Mengingat sintesis kuantitatif mencakup tiga studi, analisis sensitivitas Hartung-Knapp-Sidik-Jonkman (HKSJ) dilakukan untuk menguji ketahanan interval kepercayaan. Metode HKSJ menghasilkan estimasi titik yang sama dengan interval kepercayaan yang lebih lebar (OR = 0,164; IK 95%: 0,064 hingga 0,419; t(2) = -8,31; p = 0,014). Lebarnya interval mencerminkan kecilnya jumlah studi, namun signifikansi statistik dari efek gabungan tetap terbukti pada kedua model.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Heterogenitas: Heterogenitas statistik tidak signifikan secara bermakna (Q = 1,58, df = 2, p = 0,454; I² = 0,0%; τ² = 0,000). Bobot studi masing-masing adalah 43,5% untuk Ferdiana et al., 31,4% untuk Chalker et al., dan 25,1% untuk Onwunduba et al. Pada k = 3, uji heterogenitas memiliki power statistik yang rendah, sehingga nilai I² = 0,0% harus diartikan sebagai tidak terdeteksinya heterogenitas secara statistik, bukan bukti homogenitas mutlak (Higgins et al., 2003).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

if os.path.exists("Plots/forest_plot_core3.png"):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(6)
    p_img2.paragraph_format.space_after = Pt(4)
    run2 = p_img2.add_run()
    run2.add_picture("Plots/forest_plot_core3.png", width=Inches(6.2))
    add_p("Gambar 2. Forest Plot Meta-Analisis Model Efek Acak (Random-Effects)",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# ============================================================
# PEMBAHASAN
# ============================================================
add_p("Pembahasan", bold=True, size=12, space_before=6, space_after=6)

add_p("Temuan Utama: Meta-analisis ini, sepengetahuan kami, merupakan penelitian pertama yang menggabungkan estimasi efek secara kuantitatif khusus dari studi intervensi terkontrol di apotek komunitas swasta LMIC. Pada ketiga studi, pooled OR adalah 0,164 (IK 95%: 0,102–0,265), yang berhubungan dengan 83,6% peluang dispensing tanpa resep yang lebih rendah pada kelompok intervensi. Efek intervensi menunjukkan konsistensi di Asia Tenggara dan Afrika Sub-Sahara (I² = 0,0%), yang mengindikasikan bahwa model tata kelola berbasis apotek memiliki potensi untuk ditransfer antarkonteks.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Komparasi dengan Tinjauan Terdahulu: Tinjauan oleh Afari-Asiedu et al. (2022) dan tinjauan Cochrane (2025) tidak melakukan penggabungan data kuantitatif akibat tingginya heterogenitas desain dan setting. Namun, ketika ruang lingkup dibatasi khusus pada apotek ritel swasta yang dievaluasi menggunakan metode pasien tersimulasi, hasil studi individual menunjukkan konsistensi yang tinggi. Petugas apotek di sektor swasta menghadapi tekanan eksternal yang serupa, termasuk permintaan konsumen, kekhawatiran kehilangan pendapatan, dan keterbatasan penegakan hukum. Intervensi multifaset yang menggabungkan pelatihan staf, dukungan diagnostik cepat (seperti tes CRP), dan pengawasan sejawat terbukti mampu memediasi tekanan bisnis tersebut.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Implikasi Kebijakan: Regulasi tertulis saja tidak cukup untuk menghentikan penjualan antibiotik tanpa resep di LMIC tanpa disertai kapasitas penegakan hukum yang memadai. Hasil penelitian ini mendukung pengintegrasian paket intervensi apotek (edukasi, tes cepat, kampanye konsumen, dan peer review) ke dalam Rencana Aksi Nasional pengendalian resistensi antimikroba. Dalam kerangka Sustainable Health 5.0, stewardship apotek komunitas berperan sebagai platform intervensi terdesentralisasi yang berkelanjutan, sejalan dengan pencapaian SDG 3 (khususnya target 3.d mengenai penguatan kapasitas peringatan dini dan pengurangan risiko kesehatan global).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Keterbatasan Penelitian: Meta-analisis ini mencakup tiga studi terkontrol, mencerminkan terbatasnya uji terkontrol yang dilakukan di apotek swasta LMIC. Uji Tumwikirize et al. (2004) dari Uganda (RR 0,85; IK 95%: 0,66–1,09) tidak dapat dimasukkan dalam kuantitatif sintesis karena teks lengkap mengalami paywall dan estimasinya berupa Risk Ratio yang tidak dapat dikonversi secara andal ke aOR. Metode pasien tersimulasi mengukur perilaku dispensing pada satu waktu kunjungan tunggal dan belum merefleksikan tren konsumsi jangka panjang. Terakhir, pencarian dibatasi pada PubMed, OpenAlex, dan CENTRAL, sehingga studi regional tidak terindeks mungkin belum terekam.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ============================================================
# KESIMPULAN
# ============================================================
add_p("Kesimpulan", bold=True, size=12, space_before=6, space_after=6)

add_p("Intervensi multifaset yang menggabungkan edukasi petugas apotek, uji diagnostik cepat, kampanye kesadaran konsumen, dan pengawasan sejawat berhubungan secara signifikan dengan peluang dispensing antibiotik tanpa resep yang lebih rendah di apotek komunitas LMIC (pooled OR = 0,164; IK 95%: 0,102–0,265). Program nasional di negara berkembang disarankan untuk mengadaptasi dan mereplikasi model intervensi berbasis apotek ini dalam uji acak klaster multinegara untuk menguji keberlanjutan jangka panjang.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ============================================================
# PERNYATAAN ETIKA & DEKLARASI PENULIS
# ============================================================
add_p("Pernyataan Etika & Deklarasi Penulis", bold=True, size=12, space_before=6, space_after=6)

decl_lines = [
    "[ X ] Orisinalitas: Penulis menyatakan bahwa naskah ini adalah karya asli, belum pernah dipublikasikan sebelumnya, dan tidak sedang dalam proses pertimbangan di jurnal/konferensi lain.",
    "[ X ] Persetujuan Etika: Persetujuan komite etika tidak diperlukan karena penelitian ini merupakan tinjauan sistematis sekunder dan meta-analisis dari data sekunder yang telah dipublikasikan.",
    "[ X ] Konflik Kepentingan: Penulis menyatakan tidak memiliki konflik kepentingan finansial, personal, atau profesional terkait penelitian ini.",
    "[ X ] Persetujuan Penulis: Penulis telah meninjau, menyetujui versi naskah final ini, dan bersedia mempresentasikannya pada ICSH 2026 jika diterima.",
]
for line in decl_lines:
    add_p(line, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4)
add_p("", space_after=8)

# ============================================================
# DAFTAR PUSTAKA
# ============================================================
add_p("Daftar Pustaka", bold=True, size=12, space_before=6, space_after=6)

references_list = [
    "Murray CJL, Ikuta KS, Sharara F, Swetschinski LR, Robles Aguilar G, Gray A, et al. Global burden of bacterial antimicrobial resistance in 2019: a systematic analysis. Lancet. 2022;399(10325):629-55.",
    "Morgan DJ, Okeke IN, Laxminarayan R, Perencevich EN, Weisenberg S. Non-prescription antimicrobial use worldwide: a systematic review. Lancet Infect Dis. 2011;11(9):692-701.",
    "Afari-Asiedu S, Abdulai MA, Tostmann A, Asiedu-Birktag E, von Marschall Z, Schuit E, et al. Interventions to improve dispensing of antibiotics at the community level in low and middle income countries: a systematic review. J Glob Antimicrob Resist. 2022;29:259-74.",
    "Chalker J, Ratanawijitrasin S, Chuc NTK, Petzold M, Tomson G. Effectiveness of a multi-component intervention on dispensing practices at private pharmacies in Vietnam and Thailand: a randomized controlled trial. Soc Sci Med. 2005;60(1):131-41.",
    "Onwunduba A, Ekwunife O, Onyilogwu E, Onyemelukwe C, Modebe A, Igbokwe D. Impact of point-of-care C-reactive protein testing intervention on non-prescription dispensing of antibiotics for respiratory tract infections in private community pharmacies in Nigeria: a cluster randomized controlled trial. Int J Infect Dis. 2023;127:137-43.",
    "Ferdiana A, Wulandari LPL, Liverani M, Limato R, Essiet I, Mcknight J, et al. The impact of a multi-faceted intervention on non-prescription dispensing of antibiotics by urban community pharmacies in Indonesia: a mixed methods evaluation. BMJ Glob Health. 2024;9(10):e015620.",
    "Tumwikirize WA, Ekwaru PJ, Mohammed K, Ogwal-Okeng JW, Aupont O. Impact of a face-to-face educational intervention on improving the management of acute respiratory infections in private pharmacies and drug shops in Uganda. East Afr Med J. 2004;81(Suppl 1):S33-40.",
    "DerSimonian R, Laird N. Meta-analysis in clinical trials. Control Clin Trials. 1986;7(3):177-88.",
    "Higgins JPT, Thompson SG, Deeks JJ, Altman DG. Measuring inconsistency in meta-analyses. BMJ. 2003;327(7414):557-60."
]

for idx, ref in enumerate(references_list, 1):
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.left_indent = Inches(0.3)
    p_ref.paragraph_format.first_line_indent = Inches(-0.3)
    p_ref.paragraph_format.space_after = Pt(4)
    p_ref.paragraph_format.line_spacing = 1.15
    run = p_ref.add_run(f"{idx}.  {ref}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9.5)

doc.save(OUTPUT)
print(f"SUCCESS: Indonesian full-text manuscript saved to {OUTPUT}")
