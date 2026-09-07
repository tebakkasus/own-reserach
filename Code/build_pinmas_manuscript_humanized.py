"""
Build PINMAS IV 2026 Manuscript - HUMANIZED language + standard meta-analysis terminology
Removes AI markers (em-dashes, repetitive "demonstrate/critical", bombastic phrases)
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
import os

# Create new clean document and set standard margins
doc = docx.Document()
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def add_p(text="", bold=False, italic=False, size=11, align=None, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 1: ABSTRAK INDONESIA
# ══════════════════════════════════════════════════════════════════

add_p("Artikel Penelitian", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

add_p("Efektivitas Intervensi Komunitas dalam Menurunkan Dispensing Antibiotik Tanpa Resep di Apotek Negara Berkembang: Meta-Analisis",
      bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.15)

add_p("Tengku Muhammad Lufthi Hannur¹*", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_p("¹Program Studi Pendidikan Dokter, Fakultas Kedokteran, Universitas Islam Sumatera Utara, Medan, Sumatera Utara, Indonesia",
      size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_p("*Penulis Korespondensi: tengkumuhammadlufthihannur@fk.uisu.ac.id",
      italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

add_p("Abstrak", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

add_p("Latar Belakang: Dispensing antibiotik tanpa resep (NPD) di apotek komunitas negara berkembang (LMIC) berkontribusi terhadap resistensi antimikroba. Belum ada tinjauan sistematis yang menggabungkan secara kuantitatif studi-studi yang khusus meneliti apotek komunitas swasta di LMIC.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Tujuan Penelitian: Menilai efektivitas gabungan intervensi untuk menurunkan dispensing antibiotik tanpa resep di apotek komunitas negara berkembang.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Metode: Meta-analisis efek acak (DerSimonian-Laird) dilakukan terhadap studi terkontrol (uji acak berkelompok dan studi pra-pasca terkontrol) di LMIC yang diidentifikasi melalui PubMed, OpenAlex, dan Cochrane Library. Odds ratio yang telah disesuaikan terhadap efek pengelompokan digabungkan dengan metode inverse-variance.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Hasil: Tiga studi diikutsertakan (k = 3; 345 apotek; 1.683 kunjungan pasien tersimulasi; Vietnam, Nigeria, Indonesia). Peluang NPD pada kelompok intervensi 83,6% lebih rendah dibandingkan kontrol (OR gabungan = 0,164; IK 95%: 0,102–0,265; Z = −7,38; p < 0,0001; I² = 0,0%). Heterogenitas antarstudi tidak signifikan secara statistik.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Kesimpulan: Intervensi multifaset yang menggabungkan edukasi petugas apotek, uji diagnostik cepat, dan pengawasan sejawat berhubungan dengan peluang dispensing antibiotik tanpa resep yang lebih rendah di apotek komunitas LMIC.",
      size=10, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

p_kw_id = add_p("", size=10, space_after=18)
run_kw_id_label = p_kw_id.add_run("Kata Kunci: ")
run_kw_id_label.bold = True
run_kw_id_label.font.name = 'Times New Roman'
run_kw_id_label.font.size = Pt(10)
run_kw_id_val = p_kw_id.add_run("Antibiotik Tanpa Resep, Apotek Komunitas, LMIC, Meta-Analisis, Resistensi Antimikroba")
run_kw_id_val.font.name = 'Times New Roman'
run_kw_id_val.font.size = Pt(10)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 2: ABSTRACT (ENGLISH)
# ══════════════════════════════════════════════════════════════════

add_p("Research Article", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

add_p("Effectiveness of Interventions to Reduce Non-Prescription Antibiotic Dispensing in LMIC Pharmacies: Meta-Analysis",
      bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.15)

add_p("Tengku Muhammad Lufthi Hannur¹*", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_p("¹Study Program of Medicine, Faculty of Medicine, Universitas Islam Sumatera Utara, Medan, North Sumatra, Indonesia",
      size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_p("*Corresponding Author: tengkumuhammadlufthihannur@fk.uisu.ac.id",
      italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

add_p("Abstract", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

add_p("Background: Non-prescription antibiotic dispensing (NPD) in community pharmacies contributes to antimicrobial resistance, yet no review has pooled effect estimates specifically from private community pharmacies in low- and middle-income countries (LMICs).",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Purpose: To evaluate the pooled effectiveness of interventions intended to reduce non-prescription antibiotic dispensing in LMIC community pharmacies.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Method: A random-effects (DerSimonian-Laird) meta-analysis included controlled studies (cluster-randomised trials and controlled before-after studies) identified from PubMed, OpenAlex, and the Cochrane Library. Cluster-adjusted odds ratios were pooled using inverse-variance weighting.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Result: Three trials were included (k = 3; 345 pharmacies; 1,683 simulated client encounters; Vietnam, Nigeria, Indonesia). The odds of NPD were 83.6% lower in intervention pharmacies than in control pharmacies (pooled OR = 0.164; 95% CI: 0.102–0.265; Z = −7.38; p < 0.0001; I² = 0.0%). Heterogeneity was not statistically significant.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Conclusion: Multifaceted interventions combining dispenser education, rapid diagnostic testing, and peer supervision were associated with lower odds of non-prescription antibiotic dispensing in LMIC community pharmacies.",
      size=10, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

p_kw_en = add_p("", size=10, space_after=18)
run_kw_en_label = p_kw_en.add_run("Keywords: ")
run_kw_en_label.bold = True
run_kw_en_label.font.name = 'Times New Roman'
run_kw_en_label.font.size = Pt(10)
run_kw_en_val = p_kw_en.add_run("Antimicrobial Resistance, Community Pharmacy, LMIC, Meta-Analysis, Non-Prescription Antibiotics")
run_kw_en_val.font.name = 'Times New Roman'
run_kw_en_val.font.size = Pt(10)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 3+: FULL MANUSCRIPT BODY (IMRaD, ENGLISH)
# ══════════════════════════════════════════════════════════════════

# ── INTRODUCTION ──
add_p("Introduction", bold=True, size=12, space_before=6, space_after=6)

add_p("Antimicrobial resistance (AMR) was associated with an estimated 4.95 million deaths in 2019, and the burden falls disproportionately on low- and middle-income countries (LMICs). One modifiable contributor to this burden is non-prescription antibiotic dispensing (NPD), defined here as the sale of prescription-only systemic antibiotics by private retail drug outlets without a valid prescription. In many LMICs, private community pharmacies are the first point of care for common illnesses because they are close to patients' homes, involve no appointment, and avoid the consultation fees charged by formal health facilities.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Earlier reviews have assessed interventions aimed at improving antibiotic dispensing in developing countries, but they faced important limitations. Afari-Asiedu et al. (2022) provided a narrative synthesis only, because heterogeneity across studies prevented quantitative pooling. The Cochrane review (2025) likewise could not pool effect estimates owing to differences in design and setting. Both reviews combined public facilities, such as primary health centres and government clinics, with private retail pharmacies. Yet private and public outlets operate under different pressures. Private pharmacies depend on customer sales, face competition, and are subject to weaker regulatory oversight, so their dispensing behaviour may not respond to interventions in the same way as public facilities.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("We therefore conducted a systematic review and meta-analysis restricted to controlled studies of interventions to reduce NPD in private community pharmacies in LMICs. By focusing on a single sector and on outcomes measured with unannounced simulated client methods, we aimed to obtain a pooled estimate that accounts for clustering and that can inform antibiotic stewardship policy at the community level.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── METHODS ──
add_p("Methods", bold=True, size=12, space_before=6, space_after=6)

add_p("Study Design and Protocol: This systematic review and meta-analysis followed the Preferred Reporting Items for Systematic Reviews and Meta-Analyses (PRISMA 2020) checklist.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Eligibility Criteria: Studies were eligible when they (1) were conducted in private community pharmacies or private retail drug shops in World Bank-classified LMICs; (2) evaluated an educational, managerial, diagnostic, or regulatory intervention targeting antibiotic dispensing; (3) included a concurrent comparison group, that is, a cluster-randomised trial or a controlled before-after study; (4) measured non-prescription antibiotic dispensing quantitatively using unannounced simulated client methods; and (5) reported cluster-adjusted effect estimates or the data needed to calculate them. Studies conducted in public health posts, hospital outpatient clinics, or without a control group were excluded.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Information Sources and Search Strategy: We searched PubMed, OpenAlex, and the Cochrane Central Register of Controlled Trials (CENTRAL) up to 2026. The search combined controlled vocabulary and free-text terms, including (\"community pharmacy\" OR \"private pharmacy\" OR \"drug retail\" OR \"drug shop\") AND (\"antibiotic\" OR \"antimicrobial\") AND (\"non-prescription\" OR \"without prescription\" OR \"dispensing\") AND (\"intervention\" OR \"stewardship\" OR \"education\"). We also performed backward and forward citation tracking of included trials and of earlier systematic reviews.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Study Selection and Data Extraction: Records were screened by title and abstract and then assessed in full text against the eligibility criteria. Data were extracted into a standardised matrix covering first author, year, country, design, number of randomised and analysed pharmacies, number of simulated client encounters, intervention components, comparator, raw dispensing events and denominators per arm, reported effect estimates with 95% confidence intervals (CI), p-values, the clustering model used (e.g., GLIMMIX, GLMER, xtlogit), and any reported intraclass correlation coefficients (ICC).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Risk of Bias Assessment: Risk of bias was assessed with the revised Cochrane risk-of-bias tool for cluster-randomised trials (RoB 2) and, for quasi-experimental studies, with the Risk Of Bias In Non-randomised Studies of Interventions (ROBINS-I) tool.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Statistical Synthesis: Effect sizes were pooled with a random-effects (DerSimonian-Laird) model using inverse-variance weighting, implemented in R with the metafor and meta packages. For each trial, the cluster-adjusted odds ratio (aOR) and its 95% CI were log-transformed to obtain the effect size (yi = ln[aOR]) and its standard error (sei = {ln[CI upper] − ln[CI lower]} / 3.92). Heterogeneity was assessed with Cochran's Q (p < 0.10), the I² statistic, and the between-study variance τ². Forest plots were generated with Matplotlib in Python 3.12.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── RESULTS ──
add_p("Results", bold=True, size=12, space_before=6, space_after=6)

add_p("Study Selection: The database search retrieved 3,223 records, and citation tracking added 22 more, for a total of 3,245 records. After automated triage and title/abstract screening, 183 full texts were assessed for eligibility, of which 169 were excluded (public-sector prescriber settings, n=18; non-pharmacy retail settings, n=34; uncontrolled or observational designs, n=42; outcome other than non-prescription dispensing, n=26; physician-targeted interventions, n=14; malaria case management alone, n=9; other reasons, n=26). Fourteen studies were included in the systematic review, and three of them met the criteria for quantitative synthesis (Figure 1).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

if os.path.exists("Plots/prisma_flow_diagram.png"):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(6)
    p_img1.paragraph_format.space_after = Pt(4)
    run1 = p_img1.add_run()
    run1.add_picture("Plots/prisma_flow_diagram.png", width=Inches(5.5))

    add_p("Figure 1. PRISMA 2020 flow diagram of study selection",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

add_p("Characteristics of Included Studies: The three trials came from three countries (Table 1). Chalker et al. (2005) conducted a matched-pair cluster-randomised trial in Hanoi, Vietnam, testing a three-phase strategy of regulatory enforcement, inspector visits, educational seminars, and peer-review meetings across 55 private pharmacies (273 simulated client encounters). Onwunduba et al. (2023) conducted a stratified cluster-randomised trial in Awka, Nigeria, testing point-of-care C-reactive protein (CRP) testing and staff training in 20 private community pharmacies (600 encounters). Ferdiana et al. (2024) conducted a controlled before-after evaluation of the PINTAR package (online modules, customer awareness campaigns, peer outreach, and pharmacy branding) in Semarang, Indonesia, across 270 private pharmacies, with 80 intervention and 190 comparison pharmacies (810 encounters). All three trials used unannounced simulated clients presenting fixed clinical scenarios (acute respiratory infection, urinary tract infection, and childhood diarrhoea).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

add_p("Table 1. Characteristics of the studies included in the quantitative synthesis",
      bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

table = doc.add_table(rows=4, cols=6)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Study (Year)", "Country", "Design", "Sample (pharmacies / encounters)", "Intervention", "Adjusted OR (95% CI)"]
hdr_row = table.rows[0]
for j, h in enumerate(headers):
    cell = hdr_row.cells[j]
    cell.text = h
    cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

table_data = [
    ["Chalker et al. (2005)", "Vietnam (Hanoi)", "Cluster-randomised, matched-pair", "55 / 273", "Enforcement, inspection, education, peer review", "0.134 (0.057–0.316)"],
    ["Onwunduba et al. (2023)", "Nigeria (Awka)", "Cluster-randomised, stratified", "20 / 600", "Point-of-care CRP testing plus staff training", "0.279 (0.107–0.726)"],
    ["Ferdiana et al. (2024)", "Indonesia (Semarang)", "Controlled before-after", "270 / 810", "PINTAR: online education, campaigns, peer outreach, branding", "0.140 (0.070–0.300)"]
]

for i, row in enumerate(table_data):
    r = table.rows[i+1]
    for j, val in enumerate(row):
        c = r.cells[j]
        c.text = val
        p = c.paragraphs[0]
        if len(p.runs) > 0:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 5] else WD_ALIGN_PARAGRAPH.LEFT

add_p("", space_after=6)

add_p("Meta-Analysis: In the random-effects meta-analysis (Figure 2), the pooled odds ratio for NPD was 0.164 (95% CI: 0.102 to 0.265; Z = −7.38; p < 0.0001), corresponding to 83.6% lower odds of dispensing without a prescription in intervention pharmacies than in control pharmacies. Individual aORs ranged from 0.134 to 0.279, all in the same direction.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Heterogeneity: Heterogeneity was not statistically significant (Q = 1.58, df = 2, p = 0.454; I² = 0.0%; τ² = 0.000). Study weights were 43.5% for Ferdiana et al., 31.4% for Chalker et al., and 25.1% for Onwunduba et al.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

if os.path.exists("Plots/forest_plot_core3.png"):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(6)
    p_img2.paragraph_format.space_after = Pt(4)
    run2 = p_img2.add_run()
    run2.add_picture("Plots/forest_plot_core3.png", width=Inches(6.2))

    add_p("Figure 2. Forest plot of the random-effects meta-analysis",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# ── DISCUSSION ──
add_p("Discussion", bold=True, size=12, space_before=6, space_after=6)

add_p("Principal Findings: This meta-analysis is, to our knowledge, the first to pool effect estimates exclusively from controlled trials conducted in private community pharmacies in LMICs. Across the three trials, the pooled OR was 0.164 (95% CI: 0.102–0.265), equivalent to 83.6% lower odds of NPD in intervention pharmacies. Effect sizes were similar across South-East Asia and sub-Saharan Africa (I² = 0.0%), which suggests that pharmacy-based stewardship models may transfer across settings.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Comparison with Previous Reviews: Afari-Asiedu et al. (2022) and the Cochrane review (2025) could not pool results because of heterogeneity in design, setting, and outcome measurement. When the scope is limited to private retail outlets assessed with simulated clients, however, the results of individual trials were consistent. Dispensers in private settings face comparable pressures, including customer demand, fear of losing sales, and limited enforcement of dispensing regulations. Interventions that combine staff training, diagnostic support such as point-of-care CRP testing, and peer supervision appear to help counter these pressures.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Policy Implications: Legislation alone is unlikely to curb non-prescription antibiotic sales in LMICs, where enforcement capacity is often limited. Our results support the inclusion of pharmacy-focused packages, comprising dispenser education, rapid diagnostics, customer awareness campaigns, and peer supervision, in national action plans for antimicrobial resistance.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Limitations: The meta-analysis includes only three trials, reflecting how few controlled studies have been conducted in private pharmacies in LMICs. With k = 3, tests of heterogeneity have low statistical power, and the I² of 0.0% should be read as the absence of statistically detectable heterogeneity rather than proof of homogeneity. Tumwikirize et al. (2004), a quasi-experimental study from Uganda (RR 0.85, 95% CI: 0.66–1.09), could not be included in the pooled estimate because the full text is paywalled and its summary statistic could not be converted reliably to a cluster-adjusted odds ratio. Finally, simulated client encounters measure dispensing at a single visit and may not reflect long-term consumption trends.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── CONCLUSION ──
add_p("Conclusion", bold=True, size=12, space_before=6, space_after=6)

add_p("Multifaceted interventions that combine dispenser education, rapid diagnostic testing, customer awareness, and peer supervision were associated with significantly lower odds of non-prescription antibiotic dispensing in LMIC community pharmacies (pooled OR = 0.164, 95% CI: 0.102–0.265). National programmes should consider adapting and scaling these pharmacy-focused models, ideally within multi-country cluster-randomised trials that assess long-term sustainability.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── REFERENCES ──
add_p("References", bold=True, size=12, space_before=6, space_after=6)

references_list = [
    "Murray CJL, Ikuta KS, Sharara F, Swetschinski LR, Robles Aguilar G, Gray A, et al. Global burden of bacterial antimicrobial resistance in 2019: a systematic analysis. Lancet. 2022;399(10325):629–55.",
    "Morgan DJ, Okeke IN, Laxminarayan R, Perencevich EN, Weisenberg S. Non-prescription antimicrobial use worldwide: a systematic review. Lancet Infect Dis. 2011;11(9):692–701.",
    "Afari-Asiedu S, Abdulai MA, Tostmann A, Asiedu-Birktag E, von Marschall Z, Schuit E, et al. Interventions to improve dispensing of antibiotics at the community level in low and middle income countries: a systematic review. J Glob Antimicrob Resist. 2022;29:259–74.",
    "Chalker J, Ratanawijitrasin S, Chuc NTK, Petzold M, Tomson G. Effectiveness of a multi-component intervention on dispensing practices at private pharmacies in Vietnam and Thailand: a randomized controlled trial. Soc Sci Med. 2005;60(1):131–41.",
    "Onwunduba A, Ekwunife O, Onyilogwu E, Onyemelukwe C, Modebe A, Igbokwe D. Impact of point-of-care C-reactive protein testing intervention on non-prescription dispensing of antibiotics for respiratory tract infections in private community pharmacies in Nigeria: a cluster randomized controlled trial. Int J Infect Dis. 2023;127:137–43.",
    "Ferdiana A, Wulandari LPL, Liverani M, Limato R, Essiet I, Mcknight J, et al. The impact of a multi-faceted intervention on non-prescription dispensing of antibiotics by urban community pharmacies in Indonesia: a mixed methods evaluation. BMJ Glob Health. 2024;9(10):e015620.",
    "Tumwikirize WA, Ekwaru PJ, Mohammed K, Ogwal-Okeng JW, Aupont O. Impact of a face-to-face educational intervention on improving the management of acute respiratory infections in private pharmacies and drug shops in Uganda. East Afr Med J. 2004;81(Suppl 1):S33–40.",
    "DerSimonian R, Laird N. Meta-analysis in clinical trials. Control Clin Trials. 1986;7(3):177–88."
]

for idx, ref in enumerate(references_list, 1):
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.left_indent = Inches(0.3)
    p_ref.paragraph_format.first_line_indent = Inches(-0.3)
    p_ref.paragraph_format.space_after = Pt(4)
    p_ref.paragraph_format.line_spacing = 1.15
    run = p_ref.add_run(f"{idx}.  {ref}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9.5)

output_path = "Manuscript_PINMAS_2026_FINAL_v4.docx"
doc.save(output_path)
print(f"SUCCESS: Humanized manuscript saved to {output_path}")