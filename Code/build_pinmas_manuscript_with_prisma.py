"""
Update Manuscript PINMAS IV 2026:
1. Revise overly strong/aggressive wording (k=3 context)
2. Mention analytical software accurately (Python/R reproducible script)
3. Ensure all numbers, figures, and tables are 100% harmonized
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = docx.Document('template-abstract-pinmas-2026.docx')

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
pf = style.paragraph_format
pf.space_after = Pt(6)
pf.line_spacing = 1.5

def set_para(idx, text, bold=False, size=12, italic=False, color=None, align=None):
    p = doc.paragraphs[idx]
    p.clear()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    return p

def add_para(text, bold=False, size=12, italic=False, align=None, space_after=Pt(6)):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if align:
        p.alignment = align
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.5
    return p

# ── INDONESIAN ABSTRACT ──
set_para(0, "Artikel Penelitian", bold=True, size=12)
set_para(1,
    "Efektivitas Intervensi Komunitas dalam Menurunkan Dispensing Antibiotik "
    "Tanpa Resep di Apotek Negara Berkembang: Meta-Analisis",
    bold=True, size=14)
set_para(2, "[NAMA PENULIS]¹*", bold=False, size=12)
set_para(3,
    "¹Departemen/Program Studi, Fakultas Kedokteran/Farmasi, "
    "[Universitas/Institusi], [Kota], [Provinsi], [Negara]. "
    "*Korespondensi: [email institusi]",
    size=11)
set_para(5, "Abstrak", bold=True, size=13)
set_para(6,
    "Latar Belakang: "
    "Dispensing antibiotik tanpa resep (NPD) di apotek komunitas negara "
    "berkembang merupakan pemicu utama resistensi antimikroba global. "
    "Tinjauan sistematis terdahulu belum melakukan analisis kuantitatif "
    "yang dibatasi secara ketat pada setting apotek swasta komunitas LMIC.",
    size=12)
set_para(7,
    "Tujuan: "
    "Mengevaluasi efektivitas gabungan intervensi berorientasi apotek "
    "komunitas dalam menurunkan NPD di negara berkembang.",
    size=12)
set_para(8,
    "Metode: "
    "Meta-analisis random-effects (DerSimonian-Laird) dilakukan pada "
    "studi terkontrol (cluster RCT dan controlled pre-post) di LMIC. "
    "Efek gabungan diukur menggunakan odds ratio yang disesuaikan "
    "dengan efek pengelompokan.",
    size=12)
set_para(9,
    "Hasil: "
    "Tiga studi memenuhi kriteria inklusi (k = 3; total "
    "N = 345 apotek / 1.683 kunjungan tersimulasi; Vietnam, Nigeria, "
    "Indonesia). Intervensi menurunkan peluang NPD sebesar 83,6% "
    "(Pooled OR = 0,164; 95% CI: 0,102–0,265; Z = −7,38; p < 0,0001; "
    "I² = 0,0%).",
    size=12)
set_para(10,
    "Kesimpulan: "
    "Intervensi multifaset berbasis edukasi, diagnostik cepat, dan "
    "pengawasan edukatif terbukti berpotensi menurunkan NPD "
    "secara konsisten di apotek komunitas LMIC.",
    size=12)
set_para(12,
    "Keywords: Dispensing Antibiotik Tanpa Resep; Apotek Komunitas; LMIC; "
    "Meta-Analisis; Resistensi Antimikroba",
    size=12)

# ── ENGLISH ABSTRACT ──
set_para(14, "Research Article", bold=True, size=12)
set_para(15,
    "Effectiveness of Interventions to Reduce Non-Prescription Antibiotic "
    "Dispensing in LMIC Pharmacies: Meta-Analysis",
    bold=True, size=14)
set_para(16, "[NAMA PENULIS]¹*", bold=False, size=12)
set_para(18,
    "¹Department/Program, Faculty of Medicine/Pharmacy, "
    "[University/Institution], [City], [Province], [Country]. "
    "*Correspondence: [institutional email]",
    size=11)
set_para(20, "Abstract", bold=True, size=13)
set_para(21,
    "Background: "
    "Non-prescription antibiotic dispensing (NPD) in LMIC community "
    "pharmacies is a major driver of global antimicrobial resistance. "
    "Previous systematic reviews did not perform robust quantitative "
    "pooling restricted to private community pharmacy settings in LMICs.",
    size=12)
set_para(22,
    "Purpose: "
    "To quantitatively evaluate the pooled effectiveness of "
    "community-directed interventions in reducing NPD across LMIC "
    "private community drug retail outlets.",
    size=12)
set_para(23,
    "Method: "
    "A random-effects inverse-variance meta-analysis "
    "(DerSimonian–Laird) was conducted on controlled experimental "
    "studies (cluster RCTs and controlled pre-post) in LMICs. "
    "Pooled effect sizes used cluster-adjusted log odds ratios.",
    size=12)
set_para(24,
    "Result: "
    "Three primary studies (k = 3; N = 345 pharmacies / 1,683 "
    "simulated client encounters; Vietnam, Nigeria, Indonesia) met "
    "strict inclusion criteria. Interventions reduced odds of NPD "
    "by 83.6% (Pooled OR = 0.164; 95% CI: 0.102–0.265; "
    "Z = −7.38; p < 0.0001; I² = 0.0%).",
    size=12)
set_para(25,
    "Conclusion: "
    "Multifaceted interventions integrating dispenser education, "
    "point-of-care testing, and peer supervision show promising, "
    "consistent potential in reducing NPD across LMIC private "
    "community pharmacies.",
    size=12)
set_para(28,
    "Keywords: Antimicrobial Resistance; Community Pharmacy; LMICs; "
    "Meta-Analysis; Non-Prescription Antibiotic Dispensing",
    size=12)

# ── HEADINGS MAP ──
body_headings = ["Introduction", "Methods", "Results", "Discussion",
                 "Conclusion", "References"]
heading_idx = {}
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t in body_headings and t not in heading_idx:
        heading_idx[t] = i

# ── INTRODUCTION ──
set_para(heading_idx["Introduction"], "Introduction", bold=True, size=13)
intro_paras = [
    "Antimicrobial resistance (AMR) is a leading global health threat, "
    "with disproportionate burden in low- and middle-income countries "
    "(LMICs). Non-prescription antibiotic dispensing (NPD) — the sale "
    "of prescription-only antibiotics by private community pharmacies "
    "without a valid physician prescription — is a primary modifiable "
    "driver of AMR in these settings.",

    "Afari-Asiedu et al. (2022) conducted a systematic review of NPD "
    "interventions across LMICs but performed only narrative synthesis, "
    "as heterogeneity precluded quantitative pooling. The Cochrane "
    "review (2025) had broader global scope and similarly could not "
    "conduct meta-analysis due to methodological heterogeneity arising "
    "from the inclusion of both public-sector health posts and private "
    "retail outlets. Neither review restricted its quantitative scope "
    "specifically to private community pharmacies — the commercial "
    "sector most implicated in driving population-level AMR through "
    "unregulated dispensing practices.",

    "This study addresses this evidence gap by conducting a focused "
    "meta-analysis restricted to controlled trials evaluating "
    "interventions targeting NPD exclusively in private community "
    "pharmacies and drug retail outlets in LMICs. We aim to provide "
    "a pooled estimate of intervention effectiveness specific to the "
    "commercial dispensing context, yielding an evidence base for "
    "antimicrobial stewardship programs in community settings.",
]
for text in intro_paras:
    add_para(text)

# ── METHODS ──
set_para(heading_idx["Methods"], "Methods", bold=True, size=13)
methods_paras = [
    "Study Design and Eligibility Criteria. "
    "This systematic review and meta-analysis was conducted according "
    "to the Preferred Reporting Items for Systematic Reviews and "
    "Meta-Analyses (PRISMA) 2020 guidelines. Studies were eligible "
    "if they: (1) were conducted in private community pharmacies or "
    "retail drug outlets in World Bank-classified LMICs; (2) evaluated "
    "a structured, multicomponent intervention targeting drug dispenser "
    "practices; (3) included a concurrent control group receiving no "
    "intervention or standard care; (4) measured NPD as the primary "
    "quantitative outcome using an unannounced simulated client method "
    "(SCM) or standardised patients (SP); and (5) used a controlled "
    "experimental design (cluster randomised controlled trial or "
    "controlled pre-post quasi-experimental). Studies conducted in "
    "public-sector facilities (e.g., health centres, public clinics) were "
    "excluded.",

    "Search Strategy. "
    "A systematic search was conducted across PubMed, OpenAlex, and "
    "the Cochrane Library, supplemented by backward citation tracking "
    "of included studies and relevant systematic reviews (Afari-Asiedu "
    "et al. 2022; Cochrane 2025). The search combined terms for "
    "setting (\"community pharmacy\" OR \"drug shop\" OR \"drug "
    "retail\"), population (\"LMIC\" OR \"low-income\" OR \"developing "
    "country\"), intervention (\"intervention\" OR \"stewardship\" OR "
    "\"education\"), and outcome (\"antibiotic dispensing\" OR "
    "\"non-prescription\" OR \"without prescription\"). No date "
    "restrictions were applied.",

    "Data Extraction. "
    "Data were extracted independently by one reviewer and cross-"
    "checked against published results. Extracted variables included: "
    "study design, country, sample size (pharmacies and encounters), "
    "intervention type and duration, NPD rates (numerator/denominator) "
    "for intervention and control groups, cluster-adjusted effect "
    "sizes (odds ratios and 95% confidence intervals), statistical "
    "models used, and intraclass correlation coefficients (ICCs) where "
    "reported.",

    "Risk of Bias Assessment. "
    "Quality was assessed using the Cochrane Risk of Bias tool version "
    "2 (RoB 2) for cluster randomised trials and the Risk of Bias In "
    "Non-randomised Studies of Interventions (ROBINS-I) tool for "
    "controlled pre-post studies.",

    "Statistical Analysis. "
    "Quantitative synthesis used a random-effects model with the "
    "DerSimonian–Laird estimator applied to cluster-adjusted log "
    "odds ratios (logOR) and standard errors (SE). Heterogeneity was "
    "assessed using Cochran's Q statistic, the I² index, and τ². "
    "Publication bias was not formally assessed due to the small "
    "number of available trials (k < 10). Statistical analyses and "
    "reproducible workflows were implemented in Python 3.12 (NumPy, "
    "SciPy, Matplotlib) with cross-validation against standard meta-analytic "
    "routines in R (meta, metafor packages).",
]
for text in methods_paras:
    add_para(text)

# ── RESULTS ──
set_para(heading_idx["Results"], "Results", bold=True, size=13)
results_paras = [
    "Study Selection. "
    "A total of 3,245 records were identified from database searches "
    "(PubMed, OpenAlex, Cochrane Library; n = 3,223) and supplementary "
    "citation tracking of prior systematic reviews (n = 22). After "
    "title and abstract screening, 183 reports were assessed for "
    "full-text eligibility. Following application of strict inclusion "
    "criteria — specifically restriction to private community pharmacy "
    "settings with controlled experimental designs and simulated client "
    "outcome measures — 14 studies were retained for qualitative "
    "synthesis, and three primary studies met criteria for quantitative "
    "meta-analysis (Figure 1, PRISMA flow diagram).",
]
for text in results_paras:
    add_para(text)

# EMBED FIGURE 1 (PRISMA Flow Diagram)
prisma_img_path = "Plots/prisma_flow_diagram.png"
if os.path.exists(prisma_img_path):
    add_para("")
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p_img1.add_run()
    run1.add_picture(prisma_img_path, width=Inches(5.6))

    p_fig1_cap = add_para(
        "Figure 1. PRISMA 2020 flow diagram illustrating identification, screening, "
        "eligibility assessment, and inclusion of studies for systematic review and meta-analysis.",
        italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(12))

results_paras_2 = [
    "Characteristics of Included Studies. "
    "Three primary studies spanning three LMIC regions were synthesized: "
    "(1) Chalker et al. (2005) — a matched-pair cluster RCT in Hanoi, "
    "Vietnam (N = 55 pharmacies analyzed out of 68 randomized, 273 "
    "simulated client encounters), evaluating a multicomponent intervention "
    "combining regulatory enforcement, dispensing education, and peer review; "
    "(2) Onwunduba et al. (2023) — a stratified block cluster RCT in "
    "Awka, Nigeria (N = 20 pharmacies, 600 simulated client "
    "encounters), evaluating CRP point-of-care testing and staff "
    "training for respiratory tract infections; "
    "(3) Ferdiana et al. (2024) — a controlled pre-post quasi-"
    "experimental study (PINTAR) in Semarang, Indonesia (N = 270 "
    "pharmacies [80 participating, 190 non-participating comparison], "
    "810 standardised patient encounters), evaluating a multifaceted "
    "intervention comprising online education, customer awareness "
    "campaigns, peer pharmacy visits, and pharmacy branding. "
    "All three studies used unannounced simulated clients to assess "
    "real-world dispensing practice. Study characteristics are "
    "summarised in Table 1.",

    "Pooled Meta-Analysis. "
    "The random-effects meta-analysis yielded a pooled odds ratio of "
    "0.164 (95% CI: 0.102–0.265; Z = −7.38; p < 0.0001), indicating "
    "an 83.6% reduction in the odds of NPD in intervention versus "
    "control pharmacies (Figure 2). All individual study effect sizes "
    "were statistically significant and consistently favoured intervention. "
    "No statistically detectable heterogeneity was observed across the "
    "included trials (I² = 0.0%; τ² = 0.000; Q = 1.58; p = 0.454).",

    "Individual Study Contributions. "
    "Ferdiana et al. (2024) contributed the largest analytical weight "
    "(43.5%) with an adjusted odds ratio of 0.140 (95% CI: 0.070–"
    "0.300). Chalker et al. (2005) contributed 31.4% "
    "(OR = 0.134; 95% CI: 0.057–0.316). Onwunduba et al. (2023) "
    "contributed 25.1% (adjusted OR = 0.279; 95% CI: 0.107–0.726).",
]
for text in results_paras_2:
    add_para(text)

# TABLE 1
p_t1 = add_para("Table 1. Characteristics of Primary Studies Included in the Meta-Analysis",
                  bold=True, size=11, space_after=Pt(4))

table = doc.add_table(rows=4, cols=6)
table.style = "Table Grid"

hdr = ["Study (Year)", "Country", "Design", "N (Pharmacies /\nEncounters)",
       "Intervention", "Adj. OR (95% CI)"]
for j, h in enumerate(hdr):
    c = table.rows[0].cells[j]
    c.text = h
    for run in c.paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(9)

rows_data = [
    ["Chalker et al.\n(2005)", "Vietnam", "Cluster\nRCT",
     "55 / 273", "Regulatory enforcement +\nDispensing education +\nPeer review",
     "0.134\n(0.057–0.316)"],
    ["Onwunduba\net al. (2023)", "Nigeria", "Cluster\nRCT",
     "20 / 600", "CRP POCT kits +\nStaff training on RTI\nmanagement",
     "0.279\n(0.107–0.726)"],
    ["Ferdiana et al.\n(2024)", "Indonesia", "Controlled\nPre-Post",
     "270 / 810", "PINTAR: Education +\nCustomer campaign +\nPeer visits + Branding",
     "0.140\n(0.070–0.300)"],
]
for i, row_data in enumerate(rows_data):
    for j, val in enumerate(row_data):
        c = table.rows[i+1].cells[j]
        c.text = val
        for run in c.paragraphs[0].runs:
            run.font.size = Pt(9)

# EMBED FIGURE 2 (Forest Plot)
forest_img_path = "Plots/forest_plot_core3.png"
if os.path.exists(forest_img_path):
    add_para("")
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p_img2.add_run()
    run2.add_picture(forest_img_path, width=Inches(6.2))

    p_fig2_cap = add_para(
        "Figure 2. Forest plot of random-effects meta-analysis of interventions to "
        "reduce non-prescription antibiotic dispensing in LMIC community pharmacies. "
        "Pooled OR = 0.164 (95% CI: 0.102–0.265); I² = 0.0%.",
        italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(12))

# ── DISCUSSION ──
set_para(heading_idx["Discussion"], "Discussion", bold=True, size=13)
discussion_paras = [
    "This meta-analysis provides a focused quantitative synthesis of "
    "interventions targeting non-prescription antibiotic dispensing "
    "specifically in LMIC private community pharmacies. The pooled odds "
    "ratio of 0.164 indicates an estimated 83.6% reduction in the odds "
    "of NPD across the three included trials in Vietnam, Nigeria, and "
    "Indonesia. While previous broad reviews encountered substantial "
    "heterogeneity by combining public and private health sectors, "
    "our setting-specific focus yielded consistent point estimates "
    "and no statistically detectable heterogeneity (I² = 0.0%).",

    "This observed consistency suggests that the primary drivers of "
    "unregulated dispensing — including commercial pressure, consumer "
    "expectations, and knowledge barriers — share common structural "
    "characteristics across diverse LMIC retail pharmacy environments. "
    "Interventions that combine training, diagnostic support, and "
    "peer reinforcement appear to exert meaningful downward pressure "
    "on inappropriate dispensing across differing regional settings.",

    "Our findings corroborate and extend the narrative synthesis of "
    "Afari-Asiedu et al. (2022) by providing an initial pooled effect size "
    "restricted to the private drug retail sector. By focusing exclusively "
    "on community pharmacies evaluated via simulated client encounters, "
    "this analysis isolates the commercial retail environment where "
    "antibiotic stewardship interventions must operate against financial "
    "disincentives.",

    "Several important limitations must be acknowledged. First, the "
    "quantitative synthesis is based on a small pool of studies (k = 3), "
    "reflecting the paucity of controlled intervention trials with "
    "covert simulated client evaluations in LMIC private pharmacies. "
    "Statistical tests for heterogeneity (Q and I²) have limited power "
    "when the number of studies is very small; thus, the finding of "
    "I² = 0.0% should be interpreted cautiously rather than as definitive "
    "proof of total homogeneity. Second, Tumwikirize et al. (2004) "
    "was identified as potentially eligible from secondary review sources "
    "(reporting an RR of 0.85, 95% CI: 0.66–1.09), but was excluded "
    "from our primary OR-scale quantitative pool due to lack of direct "
    "full-text verification and non-combinable effect metrics. "
    "Third, simulated client evaluations capture immediate dispensing "
    "practices in response to specific case presentations, which may "
    "not fully capture long-term community consumption patterns.",
]
for text in discussion_paras:
    add_para(text)

# ── CONCLUSION ──
set_para(heading_idx["Conclusion"], "Conclusion", bold=True, size=13)
conclusion_paras = [
    "This meta-analysis provides quantitative evidence that multifaceted, "
    "community-directed interventions — combining dispenser education, "
    "point-of-care testing, customer engagement, and peer supervision — "
    "may substantially reduce the odds of non-prescription antibiotic "
    "dispensing in LMIC private community pharmacies.",

    "Given the small number of eligible controlled trials currently "
    "available (k = 3), these findings highlight both the promise of "
    "multifaceted retail pharmacy interventions and the urgent need "
    "for larger, multi-centre cluster randomised trials across diverse "
    "LMIC regions to evaluate the long-term sustainability and clinical "
    "impact of community pharmacy stewardship models.",
]
for text in conclusion_paras:
    add_para(text)

# ── REFERENCES ──
set_para(heading_idx["References"], "References", bold=True, size=13)
references = [
    "1. Murray CJL, Ikuta KS, Sharara F, Swetschinski LR, Robles Aguilar G, Gray A, et al. Global burden of bacterial antimicrobial resistance in 2019: a systematic analysis. Lancet. 2022;399(10325):629–55.",
    "2. Morgan DJ, Okeke IN, Laxminarayan R, Perencevich EN, Weisenberg S. Non-prescription antimicrobial use worldwide: a systematic review. Lancet Infect Dis. 2011;11(9):692–701.",
    "3. Afari-Asiedu S, Abdulai MA, Tostmann A, Asiedu-Birktag E, von Marschall Z, Schuit E, et al. Interventions to improve dispensing of antibiotics at the community level in low and middle income countries: a systematic review. J Glob Antimicrob Resist. 2022;29:259–74.",
    "4. Chalker J, Ratanawijitrasin S, Chuc NTK, Petzold M, Tomson G. Effectiveness of a multi-component intervention on dispensing practices at private pharmacies in Vietnam and Thailand—a randomized controlled trial. Soc Sci Med. 2005;60(1):131–41.",
    "5. Onwunduba A, Ekwunife O, Onyilogwu E, Onyemelukwe C, Modebe A, Igbokwe D. Impact of point-of-care C-reactive protein testing intervention on non-prescription dispensing of antibiotics for respiratory tract infections in private community pharmacies in Nigeria: a cluster randomized controlled trial. Int J Infect Dis. 2023;127:137–43.",
    "6. Ferdiana A, Wulandari LPL, Liverani M, Limato R, Essiet I, Mcknight J, et al. The impact of a multi-faceted intervention on non-prescription dispensing of antibiotics by urban community pharmacies in Indonesia: a mixed methods evaluation. BMJ Glob Health. 2024;9(10):e015620.",
    "7. Tumwikirize WA, Ekwaru PJ, Mohammed K, Ogwal-Okeng JW, Aupont O. Impact of a face-to-face educational intervention on improving the management of acute respiratory infections in private pharmacies and drug shops in Uganda. East Afr Med J. 2004;81(Suppl 1):S33–40.",
    "8. DerSimonian R, Laird N. Meta-analysis in clinical trials. Control Clin Trials. 1986;7(3):177–88.",
]

for ref in references:
    p = doc.add_paragraph()
    run = p.add_run(ref)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)

output_path = "Manuscript_PINMAS_2026_FINAL_v3.docx"
doc.save(output_path)
print(f"SUCCESS: Saved to {output_path}")
