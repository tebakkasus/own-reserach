"""
Build Full-Text Indonesian Manuscript for PINMAS IV (Teks Lengkap Bahasa Indonesia)
For study/learning and internal review.
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
import os

# Create new clean document and set standard margins
doc = docx.Document()
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Configure default font
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def add_p(text="", bold=False, italic=False, size=11, align=None, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 1: ABSTRAK (BAHASA INDONESIA)
# ══════════════════════════════════════════════════════════════════

add_p("Artikel Penelitian", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

# Judul Indonesia (14 kata)
add_p("Efektivitas Intervensi Komunitas dalam Menurunkan Dispensing Antibiotik Tanpa Resep di Apotek Negara Berkembang: Meta-Analisis",
      bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.15)

# Penulis & Afiliasi
add_p("Tengku Muhammad Lufthi Hannur¹*", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_p("¹Program Studi Pendidikan Dokter, Fakultas Kedokteran, Universitas Islam Sumatera Utara, Medan, Sumatera Utara, Indonesia",
      size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_p("*Penulis Korespondensi: tengkumuhammadlufthihannur@fk.uisu.ac.id",
      italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

# Heading Abstrak
add_p("Abstrak", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

# Isi Abstrak Terstruktur
add_p("Latar Belakang: Dispensing antibiotik tanpa resep (NPD) di apotek komunitas negara berkembang (LMIC) merupakan pemicu utama resistensi antimikroba global. Tinjauan sistematis terdahulu belum melakukan analisis kuantitatif yang dibatasi secara ketat pada setting apotek swasta komunitas LMIC.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Tujuan Penelitian: Mengevaluasi efektivitas gabungan intervensi berorientasi apotek komunitas dalam menurunkan dispensing antibiotik tanpa resep di negara berkembang.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Metode: Meta-analisis random-effects (DerSimonian-Laird) dilakukan pada studi terkontrol (cluster RCT dan controlled pre-post) di LMIC. Penelusuran sistematis dilakukan pada PubMed, OpenAlex, dan Cochrane Library. Efek gabungan diukur menggunakan log odds ratio yang disesuaikan dengan efek pengelompokan (cluster-adjusted).",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Hasil: Tiga studi memenuhi kriteria inklusi sintesis kuantitatif (k = 3; total N = 345 apotek / 1.683 kunjungan tersimulasi; Vietnam, Nigeria, Indonesia). Intervensi menurunkan peluang NPD sebesar 83,6% (Pooled OR = 0,164; 95% CI: 0,102–0,265; Z = −7,38; p < 0,0001; I² = 0,0%). Tidak ditemukan heterogenitas statistik antar-studi.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Kesimpulan: Intervensi multifaset berbasis edukasi dispenser, penyediaan uji diagnostik cepat, dan pengawasan edukatif menunjukkan potensi menurunkan dispensing antibiotik tanpa resep secara konsisten di apotek komunitas LMIC.",
      size=10, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# Kata Kunci
p_kw_id = add_p("", size=10, space_after=18)
run_kw_id_label = p_kw_id.add_run("Kata Kunci: ")
run_kw_id_label.bold = True
run_kw_id_label.font.name = 'Times New Roman'
run_kw_id_label.font.size = Pt(10)
run_kw_id_val = p_kw_id.add_run("Antibiotik Tanpa Resep, Apotek Komunitas, LMIC, Meta-Analisis, Resistensi Antimikroba")
run_kw_id_val.font.name = 'Times New Roman'
run_kw_id_val.font.size = Pt(10)

# Page Break ke Teks Lengkap
doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 2+: FULL MANUSCRIPT BODY (BAHASA INDONESIA)
# ══════════════════════════════════════════════════════════════════

# ── PENDAHULUAN ──
add_p("Pendahuluan", bold=True, size=12, space_before=6, space_after=6)

add_p("Resistensi antimikroba diperkirakan berkaitan dengan 4,95 juta kematian di dunia pada tahun 2019, dan beban terbesar ditanggung oleh negara-negara berpenghasilan rendah dan menengah (LMIC). Salah satu pendorong resistensi yang dapat dicegah adalah praktik penyerahan antibiotik tanpa resep dokter (non-prescription antibiotic dispensing/NPD). Di banyak LMIC, apotek komunitas swasta menjadi pilihan utama masyarakat saat sakit ringan karena lokasinya dekat, tidak memerlukan antrean panjang, dan tidak memungut biaya konsultasi medis formal.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Sejumlah tinjauan pustaka terdahulu telah mengevaluasi intervensi untuk memperbaiki praktik penyerahan antibiotik di negara berkembang. Namun, telaah tersebut menghadapi keterbatasan. Kajian sistematis oleh Afari-Asiedu dkk. (2022) hanya menyajikan sintesis naratif karena tingginya variasi antarstudi menghalangi penggabungan secara kuantitatif. Tinjauan Cochrane (2025) juga tidak dapat melakukan meta-analisis akibat perbedaan rancangan dan latar penelitian. Selain itu, kedua telaah tersebut menggabungkan fasilitas kesehatan pemerintah (seperti puskesmas) dengan apotek swasta. Padahal, apotek swasta beroperasi dengan dinamika yang berbeda: mereka bergantung pada penjualan harian, menghadapi persaingan dagang, dan berada di bawah pengawasan regulasi yang relatif longgar.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Oleh karena itu, penelitian ini melakukan tinjauan sistematis dan meta-analisis yang khusus dibatasi pada studi terkontrol di apotek komunitas swasta LMIC. Dengan memusatkan perhatian pada sektor swasta dan luaran yang diukur menggunakan metode kunjungan pasien tersimulasi tanpa pemberitahuan, penelitian ini bertujuan memperoleh estimasi efek gabungan yang telah disesuaikan terhadap efek pengelompokan sebagai dasar pengambilan kebijakan tata kelola antibiotik di tingkat masyarakat.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── METODE ──
add_p("Metode", bold=True, size=12, space_before=6, space_after=6)

add_p("Desain dan Protokol Studi: Tinjauan sistematis dan meta-analisis ini dilaksanakan dengan mengacu pada pedoman standar Preferred Reporting Items for Systematic Reviews and Meta-Analyses (PRISMA 2020).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Kriteria Kelayakan: Studi dimasukkan apabila memenuhi kriteria: (1) Berlokasi pada apotek komunitas swasta atau toko obat ritel swasta di negara berstatus LMIC menurut klasifikasi Bank Dunia; (2) Mengevaluasi intervensi terstruktur berupa edukasi, manajerial, sarana diagnostik cepat, atau regulasi yang ditujukan untuk menurunkan penyerahan antibiotik tanpa resep; (3) Memiliki kelompok pembanding simultan (cluster randomised controlled trial [cRCT] atau controlled pre-post quasi-experimental); (4) Mengukur outcome dispensing antibiotik tanpa resep secara kuantitatif menggunakan metode simulated client/standardised patient tanpa pemberitahuan sebelumnya; dan (5) Melaporkan estimasi efek ter-adjust klaster atau menyediakan data lengkap untuk penghitungan log odds ratio ter-adjust. Studi pada fasilitas kesehatan pemerintah/puskesmas, klinik rawat jalan rumah sakit, atau tanpa kelompok kontrol dieksklusi dari analisis kuantitatif utama.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Strategi Penelusuran Literatur: Penelusuran komprehensif dilakukan pada pangkalan data elektronik PubMed (MEDLINE), OpenAlex, dan Cochrane Central Register of Controlled Trials (CENTRAL) hingga tahun 2026. Kata kunci penelusuran mencakup kombinasi istilah MeSH dan free-text: (\"community pharmacy\" OR \"private pharmacy\" OR \"drug retail\" OR \"drug shop\") AND (\"antibiotic\" OR \"antimicrobial\" OR \"antibacterial\") AND (\"non-prescription\" OR \"without prescription\" OR \"dispensing\") AND (\"intervention\" OR \"stewardship\" OR \"education\" OR \"training\"). Penelusuran pelengkap juga dilakukan melalui backward dan forward citation tracking pada artikel yang lolos kriteria dan tinjauan sistematis sebelumnya.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Seleksi Studi dan Ekstraksi Data: Seluruh catatan hasil pencarian digabungkan dan diskrining bertahap berdasarkan judul/abstrak, dilanjutkan penilaian teks lengkap (full-text adjudication). Ekstraksi data dilakukan secara independen ke dalam matriks terstandarisasi, mencakup: nama penulis utama, tahun publikasi, negara, desain studi, jumlah apotek terandomisasi dan dianalisis, jumlah kunjungan tersimulasi, rincian komponen intervensi, karakteristik kelompok kontrol, jumlah kejadian (events) dan total kunjungan pada kedua lengan, estimasi efek yang dilaporkan (Odds Ratio, Risk Ratio, Risk Difference), 95% confidence interval, nilai p, model penyesuaian klaster (GLIMMIX, GLMER, xtlogit), serta nilai Intraclass Correlation Coefficient (ICC).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Penilaian Risiko Bias: Kualitas metodologis dan risiko bias dinilai menggunakan instrumen Cochrane Risk of Bias for Cluster-Randomised Trials (RoB 2 CRT) untuk uji acak berklaster dan Risk of Bias in Non-randomised Studies of Interventions (ROBINS-I) untuk desain kuasi-eksperimental.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Sintesis Statistik: Sintesis kuantitatif dilakukan menggunakan Generic Inverse-Variance Random-Effects Model dengan estimator DerSimonian-Laird (DL). Analisis statistik dijalankan menggunakan bahasa R melalui paket 'metafor' dan 'meta'. Nilai cluster-adjusted Odds Ratio (aOR) dan interval kepercayaan 95% ditransformasikan ke skala logaritmik untuk menghitung effect size (yi = ln(aOR)) dan standar error (sei = [ln(CI_upper) − ln(CI_lower)] / 3,92). Heterogenitas statistik diuji menggunakan uji Q Cochran (taraf signifikansi p < 0,10), indeks inkonsistensi I², dan varians antar-studi (τ²). Visualisasi Forest Plot resolusi tinggi dibuat menggunakan pustaka Matplotlib pada lingkungan Python 3.12.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── HASIL ──
add_p("Hasil", bold=True, size=12, space_before=6, space_after=6)

add_p("Seleksi Studi: Penelusuran pangkalan data menghasilkan 3.223 catatan, ditambah 22 catatan dari pelacakan sitasi sekunder, sehingga total rekaman teridentifikasi berjumlah 3.245 rekaman. Setelah proses triage otomatis dan penapisan judul/abstrak mendalam, sebanyak 183 artikel teks lengkap diperoleh untuk dievaluasi kelayakannya. Sebanyak 169 artikel dieksklusi pada tahap teks lengkap karena berlokasi di fasilitas kesehatan publik (n=18), toko non-farmasi (n=34), desain observasional tanpa kontrol (n=42), luaran non-dispensing antibiotik (n=26), intervensi berorientasi dokter puskesmas (n=14), manajemen malaria murni (n=9), dan alasan metodologis lainnya (n=26). Sebanyak 14 studi masuk ke dalam sintesis review kualitatif, dan 3 uji klinis terkontrol yang memenuhi seluruh kriteria ketat diikutsertakan dalam meta-analisis kuantitatif utama (Gambar 1).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# EMBED GAMBAR 1 (PRISMA Flow Diagram)
if os.path.exists("Plots/prisma_flow_diagram.png"):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(6)
    p_img1.paragraph_format.space_after = Pt(4)
    run1 = p_img1.add_run()
    run1.add_picture("Plots/prisma_flow_diagram.png", width=Inches(5.5))

    add_p("Gambar 1. Diagram Alir PRISMA 2020 Seleksi Studi Tinjauan Sistematis dan Meta-Analisis",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

add_p("Karakteristik Studi yang Diikutsertakan: Tiga studi inti merepresentasikan tiga kawasan LMIC yang berbeda (Tabel 1): (1) Chalker et al. (2005) di Hanoi, Vietnam (cluster RCT matched-pair) mengevaluasi intervensi multifaset 3-fase (penegakan regulasi, inspeksi berkala, seminar edukasi, dan peer-review) pada 55 apotek swasta (273 kunjungan tersimulasi); (2) Onwunduba et al. (2023) di Awka, Nigeria (cluster RCT stratified-block) mengevaluasi penyediaan kit tes cepat C-reactive protein (CRP PoCT) dan pelatihan dispenser pada 20 apotek komunitas (600 kunjungan tersimulasi); dan (3) Ferdiana et al. (2024) di Semarang, Indonesia (kuasi-eksperimental controlled pre-post/PINTAR) mengevaluasi paket multifaset (modul edukasi daring, kampanye publik, edukasi sebaya, dan branding apotek) pada 270 apotek swasta (80 apotek partisipan vs 190 apotek non-partisipan; 810 kunjungan tersimulasi). Ketiga studi mengukur praktik dispensing riil melalui kunjungan simulated client tanpa pemberitahuan dengan skenario infeksi saluran pernapasan akut (ISPA), infeksi saluran kemih (ISK), dan diare anak.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# EMBED TABEL 1
add_p("Tabel 1. Karakteristik Studi Utama yang Diikutsertakan dalam Sintesis Kuantitatif",
      bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

table = doc.add_table(rows=4, cols=6)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Penulis & Tahun", "Negara / Lokasi", "Desain Studi", "Ukuran Sampel (Apotek / Kunjungan)", "Paket Intervensi", "Adjusted OR (95% CI)"]
hdr_row = table.rows[0]
for j, h in enumerate(headers):
    cell = hdr_row.cells[j]
    cell.text = h
    cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

table_data = [
    ["Chalker et al. (2005)", "Vietnam (Hanoi)", "Cluster RCT (Matched-pair)", "55 apotek / 273 kunjungan", "Penegakan regulasi + Inspeksi + Seminar edukasi + Pertemuan peer review", "0,134 (0,057–0,316)"],
    ["Onwunduba et al. (2023)", "Nigeria (Awka)", "Cluster RCT (Stratified-block)", "20 apotek / 600 kunjungan", "Kit tes cepat CRP point-of-care + Pelatihan algoritma ISPA bagi staf", "0,279 (0,107–0,726)"],
    ["Ferdiana et al. (2024)", "Indonesia (Semarang)", "Kuasi-Eksperimental (Controlled Pre-Post)", "270 apotek / 810 kunjungan", "Program PINTAR: Modul edukasi daring + Kampanye KIE publik + Peer outreach + Branding", "0,140 (0,070–0,300)"]
]

for i, row in enumerate(table_data):
    r = table.rows[i+1]
    for j, val in enumerate(row):
        c = r.cells[j]
        c.text = val
        p = c.paragraphs[0]
        if len(p.runs) > 0:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 3, 5] else WD_ALIGN_PARAGRAPH.LEFT

add_p("", space_after=6)

add_p("Meta-Analisis Efektivitas Intervensi: Pada model random-effects meta-analisis (Gambar 2), intervensi berorientasi apotek komunitas secara signifikan menurunkan peluang dispensing antibiotik tanpa resep dibandingkan dengan kelompok kontrol (Pooled Odds Ratio = 0,164; 95% CI: 0,102–0,265; Z = −7,38; p < 0,0001). Estimasi efek gabungan ini merefleksikan penurunan peluang pemberian antibiotik tanpa resep sebesar 83,6%. Seluruh studi individual menunjukkan arah efek yang konsisten, dengan nilai aOR masing-masing berkisar antara 0,134 hingga 0,279.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Penilaian Heterogenitas: Uji statistik tidak menunjukkan adanya heterogenitas yang terdeteksi secara bermakna di antara ketiga studi (Cochran's Q = 1,58, derajat kebebasan = 2, p = 0,454; I² = 0,0%; τ² = 0,000). Bobot analitis dalam model penggabungan terdistribusi sebagai berikut: Ferdiana et al. (43,5%), Chalker et al. (31,4%), dan Onwunduba et al. (25,1%).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# EMBED GAMBAR 2 (Forest Plot)
if os.path.exists("Plots/forest_plot_core3.png"):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(6)
    p_img2.paragraph_format.space_after = Pt(4)
    run2 = p_img2.add_run()
    run2.add_picture("Plots/forest_plot_core3.png", width=Inches(6.2))

    add_p("Gambar 2. Forest Plot Random-Effects Meta-Analisis Intervensi Penurunan Dispensing Antibiotik Tanpa Resep di Apotek Komunitas LMIC",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# ── PEMBAHASAN ──
add_p("Pembahasan", bold=True, size=12, space_before=6, space_after=6)

add_p("Temuan Utama: Penelitian ini menyajikan sintesis meta-analisis kuantitatif terfokus pertama yang dibatasi secara spesifik pada ranah apotek komunitas swasta di negara berkembang. Hasil telaah menunjukkan bahwa intervensi multifaset yang terstruktur mampu menurunkan peluang penyerahan antibiotik tanpa resep secara sangat bermakna (Pooled aOR = 0,164; 95% CI: 0,102–0,265), yang setara dengan reduksi relatif sebesar 83,6%. Konsistensi arah efek di kawasan Asia Tenggara dan Afrika Sub-Sahara (I² = 0,0%) mengindikasikan potensi aplikabilitas model stewardship apotek lintas negara berkembang.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Perbandingan dengan Tinjauan Sistematis Sebelumnya: Sintesis terdahulu, seperti Afari-Asiedu et al. (2022) dan Cochrane (2025), menyimpulkan bahwa penggabungan statistik meta-analisis tidak dapat dilakukan akibat disparitas setting dan luaran. Temuan kami membuktikan bahwa ketika ruang lingkup analitis dibatasi secara homogen pada apotek ritel swasta yang dinilai dengan metode objektif simulated client, estimasi intervensi menunjukkan konsistensi yang sangat tinggi. Di sektor swasta, staf apotek menghadapi determinan struktural yang serupa—yaitu tuntutan pasien/konsumen, kekhawatiran kehilangan pendapatan finansial toko, serta pengawasan regulasi yang terbatas. Intervensi multifaset yang memadukan peningkatan kapasitas dispenser dengan sarana pendukung (seperti tes cepat CRP) dan pengawasan sebaya terbukti efektif memitigasi tekanan komersial tersebut.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Implikasi Kebijakan dan Praktik: Temuan ini menegaskan bahwa kebijakan yang semata-mata mengandalkan larangan regulasi formal dari atas ke bawah (top-down prohibition) tidak cukup efektif di LMIC karena keterbatasan kapasitas inspeksi pemerintah. Sebaliknya, adopsi strategi multifaset yang menggabungkan pelatihan berkelanjutan dispenser, integrasi perangkat diagnostik cepat di apotek, edukasi masyarakat, dan jejaring pengawasan edukatif antar-sejawat merupakan pendekatan yang sangat potensial dan realistis untuk dimasukkan ke dalam Rencana Aksi Nasional Pengendalian Resistensi Antimikroba (RAN-AMR).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Keterbatasan Studi: Beberapa keterbatasan perlu diperhatikan dalam interpretasi hasil ini. Pertama, sintesis kuantitatif utama didasarkan pada jumlah studi yang terbatas (k = 3), yang mencerminkan masih langkanya uji eksperimental terkontrol berkualitas tinggi di apotek ritel LMIC. Akibatnya, uji heterogenitas statistik (Q dan I²) memiliki kekuatan statistik (power) yang rendah, sehingga nilai I² = 0,0% harus dimaknai sebagai ketiadaan heterogenitas yang terdeteksi secara statistik dan bukan homogenitas absolut. Kedua, studi Tumwikirize et al. (2004) di Uganda yang melaporkan RR 0,85 (95% CI: 0,66–1,09) tidak dapat digabungkan ke dalam pooling model utama karena kendala akses teks lengkap berbayar dan bentuk statistik ringkasan yang berbeda. Ketiga, metode simulated client mengevaluasi perilaku penyerahan obat pada kunjungan sesaat, sehingga belum mencerminkan pola konsumsi antibiotik kumulatif jangka panjang di masyarakat.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── KESIMPULAN ──
add_p("Kesimpulan", bold=True, size=12, space_before=6, space_after=6)

add_p("Intervensi multifaset yang memadukan edukasi dispenser farmasi, penyediaan sarana diagnostik cepat, kampanye kesadaran masyarakat, dan pengawasan berbasis kelompok sebaya terbukti secara signifikan menurunkan peluang dispensing antibiotik tanpa resep di apotek komunitas negara berkembang (Pooled aOR = 0,164; 95% CI: 0,102–0,265). Otoritas kesehatan nasional dan pemangku kepentingan pengendalian AMR perlu memprioritaskan replikasi dan perluasan model intervensi berorientasi apotek komunitas ini, yang diperkuat dengan penelitian uji acak klaster multisenter berskala lebih besar guna mengevaluasi keberlanjutan efek jangka panjang.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── DAFTAR PUSTAKA ──
add_p("Daftar Pustaka", bold=True, size=12, space_before=6, space_after=6)

references_list = [
    "Murray CJL, Ikuta KS, Sharara F, Swetschinski LR, Robles Aguilar G, Gray A, et al. Global burden of bacterial antimicrobial resistance in 2019: a systematic analysis. Lancet. 2022;399(10325):629–55.",
    "Morgan DJ, Okeke IN, Laxminarayan R, Perencevich EN, Weisenberg S. Non-prescription antimicrobial use worldwide: a systematic review. Lancet Infect Dis. 2011;11(9):692–701.",
    "Afari-Asiedu S, Abdulai MA, Tostmann A, Asiedu-Birktag E, von Marschall Z,标识 Schuit E, et al. Interventions to improve dispensing of antibiotics at the community level in low and middle income countries: a systematic review. J Glob Antimicrob Resist. 2022;29:259–74.",
    "Chalker J, Ratanawijitrasin S, Chuc NTK, Petzold M, Tomson G. Effectiveness of a multi-component intervention on dispensing practices at private pharmacies in Vietnam and Thailand—a randomized controlled trial. Soc Sci Med. 2005;60(1):131–41.",
    "Onwunduba A, Ekwunife O, Onyilogwu E, Onyemelukwe C, Modebe A, Igbokwe D. Impact of point-of-care C-reactive protein testing intervention on non-prescription dispensing of antibiotics for respiratory tract infections in private community pharmacies in Nigeria: a cluster randomized controlled trial. Int J Infect Dis. 2023;127:137–43.",
    "Ferdiana A, Wulandari LPL, Liverani M, Limato R, Essiet I, Mcknight J, et al. The impact of a multi-faceted intervention on non-prescription dispensing of antibiotics by urban community pharmacies in Indonesia: a mixed methods evaluation. BMJ Glob Health. 2024;9(10):e015620.",
    "Tumwikirize WA, Ekwaru PJ, Mohammed K, Ogwal-Okeng JW, Aupont O. Impact of a face-to-face educational intervention on improving the management of acute respiratory infections in private pharmacies and drug shops in Uganda. East Afr Med J. 2004;81(Suppl 1):S33–40.",
    "DerSimonian R, Laird N. Meta-analysis in clinical trials. Control Clin Trials. 1986;7(3):177–88."
]

# Clean accidental characters in ref 3
references_list[2] = "Afari-Asiedu S, Abdulai MA, Tostmann A, Asiedu-Birktag E, von Marschall Z, Schuit E, et al. Interventions to improve dispensing of antibiotics at the community level in low and middle income countries: a systematic review. J Glob Antimicrob Resist. 2022;29:259–74."

for idx, ref in enumerate(references_list, 1):
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.left_indent = Inches(0.3)
    p_ref.paragraph_format.first_line_indent = Inches(-0.3)
    p_ref.paragraph_format.space_after = Pt(4)
    p_ref.paragraph_format.line_spacing = 1.15
    run = p_ref.add_run(f"{idx}.  {ref}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9.5)

# Save clean file
output_path = "Manuscript_PINMAS_2026_BAHASA_INDONESIA_FINAL.docx"
doc.save(output_path)
print(f"SUCCESS: Indonesian full manuscript saved to {output_path}")
