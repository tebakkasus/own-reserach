#!/usr/bin/env python3
"""
Generate ICSH_2026_Abstract_Submission_FINAL.docx
Source of truth: ICSH_2026_Abstract_Template_Originale.docx
Changes from v7:
  1. Objective section added after Background (per template P[9] explicit requirement)
  2. Track number: [TRACK NUMBER - VERIFY FROM PARTICIPANT PORTAL] (template says Select Track 01-10, cannot be guessed)
  3. Paper ID: [ORDER ID - VERIFY FROM PARTICIPANT PORTAL]
  4. Ethical Approval: adapted for secondary review (no IRB number needed)
  5. Example box (Table 1) removed per template note
Scientific content: LOCKED (no changes to statistics, results, novelty, references)
"""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE = "D:/tm/05_Own_Research/ICSH_2026_Abstract_Template_Originale.docx"
OUTPUT = "D:/tm/05_Own_Research/ICSH_2026_Abstract_Submission_FINAL.docx"

doc = docx.Document(TEMPLATE)

# ── 1. SUBMISSION METADATA TABLE (Table 0) ──
t0 = doc.tables[0].rows[0].cells[0]
t0.text = """MANUSCRIPT SUBMISSION METADATA
Important Notice for Authors: To ensure automatic metadata matching and expedited double-blind peer review, we strongly advise registering your conference account first at https://ikhconference.com/registration. Registered participants receive this template with their verified Order ID and presentation track pre-filled automatically.
Paper ID / Order ID: [ORDER ID \u2014 VERIFY FROM PARTICIPANT PORTAL]
Research Track Classification: [TRACK NUMBER \u2014 VERIFY FROM PARTICIPANT PORTAL] (Track: Medicine)
Presentation Mode: [ X ] Oral Presentation       [   ] Poster Presentation
Publication Preference: [ X ] Asian Journal of Medicine and Biomedicine (AJMB)    [   ] ICSH 2026 Proceedings"""
for p in t0.paragraphs:
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = 'Arial'; r.font.size = Pt(8.5)

# ── 2. TITLE (P4) ──
p4 = doc.paragraphs[4]
p4.text = ''
p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
p4.paragraph_format.space_before = Pt(12)
p4.paragraph_format.space_after = Pt(8)
r = p4.add_run("Effectiveness of Interventions to Reduce Non-Prescription Antibiotic Dispensing in LMIC Pharmacies: Meta-Analysis")
r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.bold = True

# ── 3. AUTHORS (P5) ──
p5 = doc.paragraphs[5]
p5.text = ''
p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
p5.paragraph_format.space_after = Pt(4)
ra = p5.add_run("Tengku Muhammad Lufthi Hannur")
ra.font.name = 'Times New Roman'; ra.font.size = Pt(11); ra.bold = True
rs = p5.add_run("1*")
rs.font.name = 'Times New Roman'; rs.font.size = Pt(10); rs.bold = True; rs.font.superscript = True

# ── 4. AFFILIATION (P6) ──
p6 = doc.paragraphs[6]
p6.text = ''
p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
p6.paragraph_format.space_after = Pt(4)
rn = p6.add_run("1 ")
rn.font.name = 'Times New Roman'; rn.font.size = Pt(8.5); rn.font.superscript = True
rf = p6.add_run("Study Program of Medicine, Faculty of Medicine, Universitas Islam Sumatera Utara, Medan, North Sumatra, Indonesia")
rf.font.name = 'Times New Roman'; rf.font.size = Pt(9.5)

# ── 5. CORRESPONDING AUTHOR (P7) ──
p7 = doc.paragraphs[7]
p7.text = ''
p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
p7.paragraph_format.space_after = Pt(12)
rc = p7.add_run("* Corresponding Author: Tengku Muhammad Lufthi Hannur | Email: tengkumuhammadlufthihannur@fk.uisu.ac.id | WhatsApp: +62 822-6740-3241 | ORCID: https://orcid.org/0009-0008-1306-2485")
rc.font.name = 'Times New Roman'; rc.font.size = Pt(9.0); rc.italic = True

# ── 6. ABSTRACT BODY (P9) — Background, Objective, Methods, Results, Conclusion ──
p9 = doc.paragraphs[9]
p9.text = ''
p9.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p9.paragraph_format.line_spacing = 1.15
p9.paragraph_format.space_after = Pt(8)

abstract_sections = [
    ("Background: ", "Non-prescription antibiotic dispensing (NPD) in community pharmacies contributes to antimicrobial resistance, yet no quantitative review has pooled intervention effect estimates specifically from private community pharmacies in low- and middle-income countries (LMICs). "),
    ("Objective: ", "To evaluate the pooled effectiveness of interventions intended to reduce non-prescription antibiotic dispensing in private community pharmacies in LMICs. "),
    ("Methods: ", "We conducted a systematic review and meta-analysis following PRISMA 2020 guidelines. PubMed, OpenAlex, and the Cochrane Central Register of Controlled Trials were searched up to September 2026. Eligible studies were controlled intervention studies evaluating interventions to reduce NPD in private community pharmacies in LMICs, with outcomes assessed by unannounced simulated client methods. Cluster-adjusted odds ratios were pooled using a random-effects model with the DerSimonian-Laird estimator and inverse-variance weighting. A Hartung-Knapp-Sidik-Jonkman sensitivity analysis was performed to assess the robustness of the confidence interval given the small number of included studies. "),
    ("Results: ", "Three controlled studies from Vietnam, Nigeria, and Indonesia met eligibility criteria (k = 3; 345 pharmacies; 1,683 simulated client encounters). The odds of NPD were 83.6% lower in intervention pharmacies compared with control pharmacies (pooled OR = 0.164; 95% CI: 0.102 to 0.265; Z = -7.38; p < 0.0001). The Hartung-Knapp-Sidik-Jonkman method yielded a wider interval (OR = 0.164; 95% CI: 0.064 to 0.419; t(2) = -8.31, p = 0.014). Statistical heterogeneity was not detected (I-squared = 0.0%; Q = 1.58, df = 2, p = 0.454). "),
    ("Conclusion: ", "To our knowledge, this is the first meta-analysis to pool controlled intervention data on NPD interventions from private community pharmacies in LMICs. Multifaceted interventions combining dispenser education, rapid diagnostic testing, and peer supervision were associated with lower odds of non-prescription antibiotic dispensing.")
]

for label, text in abstract_sections:
    rl = p9.add_run(label); rl.font.name = 'Times New Roman'; rl.font.size = Pt(10.5); rl.bold = True
    rt = p9.add_run(text);  rt.font.name = 'Times New Roman'; rt.font.size = Pt(10.5)

# ── 7. KEYWORDS (P10) ──
p10 = doc.paragraphs[10]
p10.text = ''
p10.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p10.paragraph_format.space_after = Pt(14)
rkl = p10.add_run("Keywords: "); rkl.font.name='Times New Roman'; rkl.font.size=Pt(9.5); rkl.bold=True
rkv = p10.add_run("Antimicrobial Resistance; Community Pharmacy; Non-Prescription Antibiotic Dispensing; Meta-Analysis; Low- and Middle-Income Countries")
rkv.font.name='Times New Roman'; rkv.font.size=Pt(9.5)

# ── 8. ETHICAL DECLARATION (P13) ──
p13 = doc.paragraphs[13]
p13.text = ''
p13.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p13.paragraph_format.line_spacing = 1.15
rd = p13.add_run(
    "[ X ] Originality: The authors certify that this abstract is original, has not been published previously, and is not under consideration elsewhere.\n"
    "[ X ] Ethical Approval: Ethical approval was not required as this study is a secondary systematic review and meta-analysis of publicly available, published trial data. No primary human subjects were involved.\n"
    "[ X ] Conflict of Interest: The authors declare that they have no financial, personal, or professional conflicts of interest regarding this work.\n"
    "[ X ] Author Consent: The listed author has reviewed, approved this submission, and agreed to present at ICSH 2026 if accepted."
)
rd.font.name='Times New Roman'; rd.font.size=Pt(8.5)

# ── 9. REMOVE EXAMPLE BOX (Table 1) per template instruction ──
doc.tables[1]._element.getparent().remove(doc.tables[1]._element)

doc.save(OUTPUT)
print(f"SUCCESS: {OUTPUT}")
