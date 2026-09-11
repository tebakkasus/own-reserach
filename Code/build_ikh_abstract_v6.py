#!/usr/bin/env python3
"""
Generate ICSH 2026 Official Abstract Submission Document (v6 Final).
Applies exact surgical patches:
  1. HKSJ p-value fixed to exact mathematical calculation: p = 0.014 (t(2) = -8.31)
  2. Abstract design wording: "Eligible studies were controlled intervention studies..."
  3. Abstract structure verified: Background / Methods / Results / Conclusion (no spurious Discussion)
  4. Ethical declaration checklist strictly aligned with template facts
  5. Submission metadata table embedded per template with Track: Medicine (Sustainable Health 5.0)
  6. Author block fully populated with real contact details
"""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_PATH = "D:/tm/05_Own_Research/ICSH_2026_Abstract_Template_Original.docx"
OUTPUT_PATH = "D:/tm/05_Own_Research/ICSH_2026_Abstract_Submission_v6.docx"

doc = docx.Document(TEMPLATE_PATH)

# ── 1. SUBMISSION METADATA TABLE (Table 0) ──
t0 = doc.tables[0].rows[0].cells[0]
t0.text = """MANUSCRIPT SUBMISSION METADATA
Important Notice for Authors: To ensure automatic metadata matching and expedited double-blind peer review, we strongly advise registering your conference account first at https://ikhconference.com/registration. Registered participants receive this template with their verified Order ID and presentation track pre-filled automatically.
Paper ID / Order ID: [ PENDING REGISTRATION ]
Research Track Classification: Track 02: Medicine (Sustainable Health 5.0)
Presentation Mode: [ X ] Oral Presentation       [   ] Poster Presentation
Publication Preference: [ X ] Asian Journal of Medicine and Biomedicine (AJMB)    [   ] ICSH 2026 Proceedings"""

for p in t0.paragraphs:
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)

# ── 2. TITLE (P4) ──
p4 = doc.paragraphs[4]
p4.text = ''
p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
p4.paragraph_format.space_before = Pt(12)
p4.paragraph_format.space_after = Pt(8)
r_title = p4.add_run("Effectiveness of Interventions to Reduce Non-Prescription Antibiotic Dispensing in LMIC Pharmacies: Meta-Analysis")
r_title.font.name = 'Times New Roman'
r_title.font.size = Pt(14)
r_title.bold = True

# ── 3. AUTHORS (P5) ──
p5 = doc.paragraphs[5]
p5.text = ''
p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
p5.paragraph_format.space_after = Pt(4)
r_auth = p5.add_run("Tengku Muhammad Lufthi Hannur")
r_auth.font.name = 'Times New Roman'
r_auth.font.size = Pt(11)
r_auth.bold = True
r_sup = p5.add_run("1*")
r_sup.font.name = 'Times New Roman'
r_sup.font.size = Pt(10)
r_sup.bold = True
r_sup.font.superscript = True

# ── 4. AFFILIATION (P6) ──
p6 = doc.paragraphs[6]
p6.text = ''
p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
p6.paragraph_format.space_after = Pt(4)
r_aff_num = p6.add_run("1 ")
r_aff_num.font.name = 'Times New Roman'
r_aff_num.font.size = Pt(8.5)
r_aff_num.font.superscript = True
r_aff = p6.add_run("Study Program of Medicine, Faculty of Medicine, Universitas Islam Sumatera Utara, Medan, North Sumatra, Indonesia")
r_aff.font.name = 'Times New Roman'
r_aff.font.size = Pt(9.5)

# ── 5. CORRESPONDING AUTHOR (P7) ──
p7 = doc.paragraphs[7]
p7.text = ''
p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
p7.paragraph_format.space_after = Pt(12)
r_corr = p7.add_run("* Corresponding Author: Tengku Muhammad Lufthi Hannur | Email: tengkumuhammadlufthihannur@fk.uisu.ac.id | WhatsApp: +62 822-6740-3241 | ORCID: https://orcid.org/0009-0008-1306-2485")
r_corr.font.name = 'Times New Roman'
r_corr.font.size = Pt(9.0)
r_corr.italic = True

# ── 6. ABSTRACT BODY (P9) ──
p9 = doc.paragraphs[9]
p9.text = ''
p9.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p9.paragraph_format.line_spacing = 1.15
p9.paragraph_format.space_after = Pt(8)

abstract_sections = [
    ("Background: ", "Non-prescription antibiotic dispensing (NPD) in community pharmacies contributes to antimicrobial resistance, yet no quantitative review has pooled intervention effect estimates specifically from private community pharmacies in low- and middle-income countries (LMICs). "),
    ("Methods: ", "We conducted a systematic review and meta-analysis following PRISMA 2020 guidelines. PubMed, OpenAlex, and the Cochrane Central Register of Controlled Trials were searched up to September 2026. Eligible studies were controlled intervention studies evaluating interventions to reduce NPD in private community pharmacies in LMICs, with outcomes assessed by unannounced simulated client methods. Cluster-adjusted odds ratios were pooled using a random-effects model with the DerSimonian-Laird estimator and inverse-variance weighting. A Hartung-Knapp-Sidik-Jonkman sensitivity analysis was performed to assess the robustness of the confidence interval given the small number of included studies. "),
    ("Results: ", "Three trials from Vietnam, Nigeria, and Indonesia met eligibility criteria (k = 3; 345 pharmacies; 1,683 simulated client encounters). The odds of NPD were 83.6% lower in intervention pharmacies compared with control pharmacies (pooled OR = 0.164; 95% CI: 0.102 to 0.265; Z = -7.38; p < 0.0001). The Hartung-Knapp-Sidik-Jonkman method yielded a wider interval (OR = 0.164; 95% CI: 0.064 to 0.419; t(2) = -8.31, p = 0.014). Statistical heterogeneity was not detected (I-squared = 0.0%; Q = 1.58, df = 2, p = 0.454). "),
    ("Conclusion: ", "To our knowledge, this is the first meta-analysis to pool controlled trial data on NPD interventions from private community pharmacies in LMICs. Multifaceted interventions combining dispenser education, rapid diagnostic testing, and peer supervision were associated with lower odds of non-prescription antibiotic dispensing.")
]

for label, text in abstract_sections:
    rl = p9.add_run(label); rl.font.name = 'Times New Roman'; rl.font.size = Pt(10.5); rl.bold = True
    rt = p9.add_run(text);  rt.font.name = 'Times New Roman'; rt.font.size = Pt(10.5)

# ── 7. KEYWORDS (P10) ──
p10 = doc.paragraphs[10]
p10.text = ''
p10.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p10.paragraph_format.space_after = Pt(14)
r_kw_lbl = p10.add_run("Keywords: "); r_kw_lbl.font.name = 'Times New Roman'; r_kw_lbl.font.size = Pt(9.5); r_kw_lbl.bold = True
r_kw_val = p10.add_run("Antimicrobial Resistance; Community Pharmacy; Non-Prescription Antibiotic Dispensing; Meta-Analysis; Low- and Middle-Income Countries")
r_kw_val.font.name = 'Times New Roman'; r_kw_val.font.size = Pt(9.5)

# ── 8. ETHICAL DECLARATION (P13) ──
p13 = doc.paragraphs[13]
p13.text = ''
p13.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p13.paragraph_format.line_spacing = 1.15

decl_text = """[ X ] Originality: The authors certify that this abstract is original, has not been published previously, and is not under consideration elsewhere.
[ X ] Ethical Approval: Ethical approval was not required as this study is a secondary systematic review and meta-analysis of publicly available, published trial data.
[ X ] Conflict of Interest: The authors declare that they have no financial, personal, or professional conflicts of interest regarding this work.
[ X ] Author Consent: The listed author has reviewed, approved this submission, and agreed to present at ICSH 2026 if accepted."""

r_decl = p13.add_run(decl_text)
r_decl.font.name = 'Times New Roman'
r_decl.font.size = Pt(8.5)

# ── 9. REMOVE REFERENCE BOX (Table 1) ──
table1_elem = doc.tables[1]._element
table1_elem.getparent().remove(table1_elem)

doc.save(OUTPUT_PATH)
print(f"SUCCESS: Generated ICSH 2026 Abstract Submission v6 at {OUTPUT_PATH}")
