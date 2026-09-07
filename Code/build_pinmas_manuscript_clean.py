"""
Build 100% Clean PINMAS IV 2026 Manuscript
Completely wipes all template instructional/dummy text and builds a publication-ready document.
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
import os

# Open the base template to retain page margins, headers, footers, and styles
doc = docx.Document('template-abstract-pinmas-2026.docx')

# Remove ALL existing tables (including dummy template tables)
for table in list(doc.tables):
    p = table._element.getparent()
    p.remove(table._element)

# Clear ALL existing paragraphs
# In python-docx, we can clear text or rebuild body
for p in list(doc.paragraphs):
    p_elem = p._element
    p_parent = p_elem.getparent()
    p_parent.remove(p_elem)

# Re-configure default font
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def add_p(text="", bold=False, italic=False, size=11, align=None, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 1: ABSTRAK INDONESIA
# ══════════════════════════════════════════════════════════════════

add_p("Artikel Penelitian", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

# Judul Indonesia (14 kata, Capital Each Word kecuali kata hubung)
add_p("Efektivitas Intervensi Komunitas dalam Menurunkan Dispensing Antibiotik Tanpa Resep di Apotek Negara Berkembang: Meta-Analisis",
      bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.15)

# Penulis & Afiliasi
add_p("Tengku Muhammad Lufthi Hannur¹*", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_p("¹Program Studi Pendidikan Dokter, Fakultas Kedokteran, Universitas Islam Sumatera Utara, Medan, Sumatera Utara, Indonesia",
      size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_p("*Penulis Korespondensi: tengkumuhammadlufthihannur@fk.uisu.ac.id",
      italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

# Heading Abstrak
add_p("Abstrak", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

# Isi Abstrak Terstruktur
add_p("Latar Belakang: Dispensing antibiotik tanpa resep (NPD) di apotek komunitas negara berkembang (LMIC) merupakan pemicu utama resistensi antimikroba global. Tinjauan sistematis terdahulu belum melakukan analisis kuantitatif yang dibatasi secara ketat pada setting apotek swasta komunitas LMIC.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Tujuan Penelitian: Mengevaluasi efektivitas gabungan intervensi berorientasi apotek komunitas dalam menurunkan dispensing antibiotik tanpa resep di negara berkembang.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Metode: Meta-analisis random-effects (DerSimonian-Laird) dilakukan pada studi terkontrol (cluster RCT dan controlled pre-post) di LMIC. Penelusuran sistematis dilakukan pada PubMed, OpenAlex, dan Cochrane Library. Efek gabungan diukur menggunakan log odds ratio yang disesuaikan dengan efek pengelompokan (cluster-adjusted).",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Hasil: Tiga studi memenuhi kriteria inklusi sintesis kuantitatif (k = 3; total N = 345 apotek / 1.683 kunjungan tersimulasi; Vietnam, Nigeria, Indonesia). Intervensi menurunkan peluang NPD sebesar 83,6% (Pooled OR = 0,164; 95% CI: 0,102–0,265; Z = −7,38; p < 0,0001; I² = 0,0%). Tidak ditemukan heterogenitas statistik antar-studi.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Kesimpulan: Intervensi multifaset berbasis edukasi dispenser, penyediaan uji diagnostik cepat, dan pengawasan edukatif menunjukkan potensi menurunkan dispensing antibiotik tanpa resep secara konsisten di apotek komunitas LMIC.",
      size=10, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# Kata Kunci
p_kw_id = add_p("", size=10, space_after=18)
run_kw_id_label = p_kw_id.add_run("Kata Kunci: ")
run_kw_id_label.bold = True
run_kw_id_label.font.name = 'Times New Roman'
run_kw_id_label.font.size = Pt(10)
run_kw_id_val = p_kw_id.add_run("Antibiotik Tanpa Resep, Apotek Komunitas, LMIC, Meta-Analisis, Resistensi Antimikroba")
run_kw_id_val.font.name = 'Times New Roman'
run_kw_id_val.font.size = Pt(10)

# Page Break ke Abstrak Bahasa Inggris
doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 2: ABSTRACT (ENGLISH)
# ══════════════════════════════════════════════════════════════════

add_p("Research Article", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

# Judul Inggris (12 kata)
add_p("Effectiveness of Interventions to Reduce Non-Prescription Antibiotic Dispensing in LMIC Pharmacies: Meta-Analysis",
      bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.15)

# Author & Affiliation
add_p("Tengku Muhammad Lufthi Hannur¹*", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_p("¹Study Program of Medicine, Faculty of Medicine, Universitas Islam Sumatera Utara, Medan, North Sumatra, Indonesia",
      size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_p("*Corresponding Author: tengkumuhammadlufthihannur@fk.uisu.ac.id",
      italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

# Heading Abstract
add_p("Abstract", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

# Structured Abstract Body
add_p("Background: Non-prescription antibiotic dispensing (NPD) in low- and middle-income country (LMIC) community pharmacies is a critical driver of global antimicrobial resistance. Previous systematic reviews have not performed focused quantitative pooling restricted specifically to private community pharmacy settings in LMICs.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Purpose: To quantitatively evaluate the pooled effectiveness of community pharmacy-directed interventions in reducing non-prescription antibiotic dispensing across LMIC retail drug outlets.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Method: A random-effects inverse-variance meta-analysis (DerSimonian-Laird) was conducted on controlled experimental studies (cluster RCTs and controlled pre-post) in LMICs identified from PubMed, OpenAlex, and Cochrane Library. Pooled effect sizes were calculated using cluster-adjusted log odds ratios (logOR).",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Result: Three primary trials (k = 3; total N = 345 pharmacies / 1,683 simulated client encounters; Vietnam, Nigeria, Indonesia) met strict inclusion criteria. Interventions significantly reduced the odds of NPD by 83.6% (Pooled OR = 0.164; 95% CI: 0.102–0.265; Z = −7.38; p < 0.0001; I² = 0.0%). No statistically detectable heterogeneity was observed across studies.",
      size=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_p("Conclusion: Multifaceted interventions integrating dispenser education, rapid diagnostic testing, and peer supervision show promising potential to consistently reduce non-prescription antibiotic dispensing in LMIC community pharmacies.",
      size=10, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# Keywords
p_kw_en = add_p("", size=10, space_after=18)
run_kw_en_label = p_kw_en.add_run("Keywords: ")
run_kw_en_label.bold = True
run_kw_en_label.font.name = 'Times New Roman'
run_kw_en_label.font.size = Pt(10)
run_kw_en_val = p_kw_en.add_run("Antimicrobial Resistance, Community Pharmacy, LMIC, Meta-Analysis, Non-Prescription Antibiotics")
run_kw_en_val.font.name = 'Times New Roman'
run_kw_en_val.font.size = Pt(10)

# Page Break ke Body Naskah
doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
#  HALAMAN 3+: FULL MANUSCRIPT BODY (IMRaD)
# ══════════════════════════════════════════════════════════════════

# ── INTRODUCTION ──
add_p("Introduction", bold=True, size=12, space_before=6, space_after=6)

add_p("Antimicrobial resistance (AMR) represents one of the most critical threats to global health, causing an estimated 4.95 million associated deaths annually, with disproportionate morbidity in low- and middle-income countries (LMICs). A major modifiable driver of AMR in these settings is non-prescription antibiotic dispensing (NPD) — the direct commercial sale of prescription-only systemic antibiotics by private retail drug sellers without a prescription from an authorized prescriber. In many LMICs, private community pharmacies serve as the primary first point of healthcare contact due to geographic convenience, shorter waiting times, and avoidance of consultation fees in formal facilities.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Previous systematic reviews have examined interventions to improve antibiotic dispensing practices in developing countries. However, existing reviews suffer from critical methodological limitations. The systematic review by Afari-Asiedu et al. (2022) provided only a narrative synthesis because substantial heterogeneity across included studies precluded quantitative pooling. Similarly, the Cochrane review (2025) evaluated interventions across global health settings but was unable to pool effect estimates due to high methodological and setting diversity. Importantly, these reviews combined public health facilities (such as primary health centers and government clinics) with private retail pharmacies. Conflating public and private healthcare contexts obscures the commercial profit incentives, customer bargaining dynamics, and regulatory environments unique to private community pharmacies.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("To address this knowledge gap, the present study performs a focused systematic review and meta-analysis strictly restricted to controlled experimental studies evaluating interventions targeting non-prescription antibiotic dispensing in private community pharmacies in LMICs. By isolating the private drug retail sector and standardizing outcomes measured via unannounced simulated client methods, this study aims to deliver a precise, cluster-adjusted pooled effect estimate to inform community-level antimicrobial stewardship policies.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── METHODS ──
add_p("Methods", bold=True, size=12, space_before=6, space_after=6)

add_p("Study Design and Protocol: This systematic review and meta-analysis was conducted in accordance with the Preferred Reporting Items for Systematic Reviews and Meta-Analyses (PRISMA 2020) guidelines.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Eligibility Criteria: Studies were eligible for inclusion if they: (1) were conducted in private community pharmacies or private retail drug shops in World Bank-classified LMICs; (2) evaluated a structured educational, managerial, diagnostic, or regulatory intervention targeting antibiotic dispensing practices; (3) included a concurrent comparison group (cluster randomised controlled trial [cRCT] or controlled pre-post quasi-experimental design); (4) measured non-prescription antibiotic dispensing as a primary or secondary quantitative outcome using unannounced Simulated Client Methods (SCM) or Standardised Patients (SP); and (5) reported cluster-adjusted effect estimates or sufficient data to calculate adjusted effect sizes. Studies conducted in public sector health posts, hospital outpatient clinics, or without a concurrent control group were excluded.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Information Sources and Search Strategy: A comprehensive systematic search was performed across PubMed, OpenAlex, and the Cochrane Central Register of Controlled Trials (CENTRAL) up to 2026. The search strategy combined controlled vocabulary (MeSH) and free-text terms: (\"community pharmacy\" OR \"private pharmacy\" OR \"drug retail\" OR \"drug shop\") AND (\"antibiotic\" OR \"antimicrobial\" OR \"antibacterial\") AND (\"non-prescription\" OR \"without prescription\" OR \"dispensing\") AND (\"intervention\" OR \"stewardship\" OR \"education\" OR \"training\"). In addition, backward and forward citation tracking of included trials and prior systematic reviews was conducted.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Study Selection and Data Extraction: Deduplicated search records were systematically screened by title and abstract, followed by full-text adjudication against predefined eligibility criteria. Data extraction was conducted using a standardized data extraction matrix, capturing: first author, publication year, country, study design, number of randomized and analyzed pharmacies, number of simulated client encounters, intervention components and duration, comparator description, raw dispensing events and denominators per study arm, reported effect sizes (Odds Ratio, Risk Ratio, Risk Difference), 95% confidence intervals, p-values, cluster-adjustment models (e.g., GLIMMIX, GLMER, XTLOGIT), and intraclass correlation coefficients (ICC).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Risk of Bias Assessment: Methodological quality and risk of bias were assessed using the Cochrane Risk of Bias tool for cluster-randomised trials (RoB 2 CRT) for randomised trials and the Risk of Bias in Non-randomised Studies of Interventions (ROBINS-I) tool for quasi-experimental designs.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Statistical Synthesis: Quantitative synthesis was conducted using a Generic Inverse-Variance Random-Effects Model with the DerSimonian-Laird (DL) estimator. Statistical analyses were performed in R using the metafor and meta packages. Cluster-adjusted Odds Ratios (aOR) and their 95% confidence intervals were log-transformed to derive effect sizes (yi = ln(aOR)) and standard errors (sei = (ln(CI_upper) − ln(CI_lower)) / 3.92). Statistical heterogeneity was assessed using Cochran's Q statistic (significance threshold p < 0.10), the I² inconsistency index, and between-study variance (τ²). High-resolution forest plot visualizations were generated using Python 3.12 (Matplotlib).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── RESULTS ──
add_p("Results", bold=True, size=12, space_before=6, space_after=6)

add_p("Study Selection: The database search retrieved 3,223 records, supplemented by 22 records identified through backward citation tracking of prior systematic reviews, yielding a total pool of 3,245 records. Following automated triage and title/abstract screening, 183 full-text articles were retrieved and assessed for eligibility. A total of 169 reports were excluded during full-text review due to non-private settings (public health facilities, n=18), non-pharmacy retail environments (n=34), non-controlled observational designs (n=42), irrelevant primary outcomes (n=26), primary care physician-targeted interventions (n=14), malaria-only case management (n=9), and other non-qualifying features (n=26). Fourteen studies were included in the qualitative review pool, from which three primary controlled trials meeting all strict criteria were included in the primary quantitative meta-analysis (Figure 1).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# EMBED FIGURE 1 (PRISMA Flow Diagram)
if os.path.exists("Plots/prisma_flow_diagram.png"):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(6)
    p_img1.paragraph_format.space_after = Pt(4)
    run1 = p_img1.add_run()
    run1.add_picture("Plots/prisma_flow_diagram.png", width=Inches(5.5))

    add_p("Figure 1. PRISMA 2020 Flow Diagram of Study Selection for Systematic Review and Meta-Analysis",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

add_p("Characteristics of Included Studies: The three core trials spanned three distinct LMIC regions (Table 1): (1) Chalker et al. (2005) conducted a matched-pair cluster RCT in Hanoi, Vietnam, evaluating a 3-phase multicomponent strategy (regulatory enforcement, inspector visits, educational seminars, and peer-review meetings) across 55 private pharmacies (273 simulated client visits); (2) Onwunduba et al. (2023) conducted a stratified-block cluster RCT in Awka, Nigeria, evaluating point-of-care C-reactive protein (CRP) test kits and staff training across 20 private community pharmacies (600 simulated client visits); and (3) Ferdiana et al. (2024) conducted a controlled pre-post quasi-experimental trial (PINTAR) in Semarang, Indonesia, evaluating a multifaceted package (online interactive modules, customer awareness campaigns, peer outreach, and pharmacy branding) across 270 private pharmacies (80 participating vs 190 non-participating comparison pharmacies; 810 standardized patient visits). All three studies measured actual dispensing behavior using unannounced simulated client scenarios (acute respiratory infection, urinary tract infection, and child diarrhea).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# EMBED TABLE 1
add_p("Table 1. Characteristics of Core Studies Included in the Quantitative Synthesis",
      bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

table = doc.add_table(rows=4, cols=6)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Study & Year", "Country", "Study Design", "Sample Size (Pharmacies / Visits)", "Intervention Package", "Adjusted OR (95% CI)"]
hdr_row = table.rows[0]
for j, h in enumerate(headers):
    cell = hdr_row.cells[j]
    cell.text = h
    cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

table_data = [
    ["Chalker et al. (2005)", "Vietnam (Hanoi)", "Cluster RCT (Matched-pair)", "55 pharmacies / 273 encounters", "Regulatory enforcement + Inspector visits + Education + Peer review", "0.134 (0.057–0.316)"],
    ["Onwunduba et al. (2023)", "Nigeria (Awka)", "Cluster RCT (Stratified-block)", "20 PCPs / 600 encounters", "Semi-quantitative CRP point-of-care test kits + Staff training on RTI algorithms", "0.279 (0.107–0.726)"],
    ["Ferdiana et al. (2024)", "Indonesia (Semarang)", "Quasi-experimental (Controlled Pre-Post)", "270 pharmacies / 810 encounters", "PINTAR: Online education + Customer IEC campaign + Peer visits + Branding", "0.140 (0.070–0.300)"]
]

for i, row in enumerate(table_data):
    r = table.rows[i+1]
    for j, val in enumerate(row):
        c = r.cells[j]
        c.text = val
        p = c.paragraphs[0]
        if len(p.runs) > 0:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 3, 5] else WD_ALIGN_PARAGRAPH.LEFT

add_p("", space_after=6)

add_p("Meta-Analysis of Intervention Effectiveness: In the random-effects meta-analysis (Figure 2), community pharmacy-directed interventions demonstrated a marked and statistically significant reduction in the odds of non-prescription antibiotic dispensing compared to control pharmacies (Pooled Odds Ratio = 0.164, 95% CI: 0.102 to 0.265, Z = −7.38, p < 0.0001). This pooled estimate reflects an overall 83.6% reduction in the odds of inappropriate antibiotic dispensing. All three included trials demonstrated individual effect estimates in the same direction, with individual adjusted ORs ranging from 0.134 to 0.279.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Assessment of Heterogeneity: Statistical testing revealed an absence of detectable heterogeneity across the three trials (Cochran's Q = 1.58, degrees of freedom = 2, p = 0.454; I² = 0.0%; τ² = 0.000). Analytical weights in the pooling model were distributed across studies: Ferdiana et al. (43.5%), Chalker et al. (31.4%), and Onwunduba et al. (25.1%).",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# EMBED FIGURE 2 (Forest Plot)
if os.path.exists("Plots/forest_plot_core3.png"):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(6)
    p_img2.paragraph_format.space_after = Pt(4)
    run2 = p_img2.add_run()
    run2.add_picture("Plots/forest_plot_core3.png", width=Inches(6.2))

    add_p("Figure 2. Forest Plot of Random-Effects Meta-Analysis of Interventions to Reduce Non-Prescription Antibiotic Dispensing in LMIC Community Pharmacies",
          bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# ── DISCUSSION ──
add_p("Discussion", bold=True, size=12, space_before=6, space_after=6)

add_p("Principal Findings: This study provides the first focused quantitative meta-analysis restricted exclusively to private community pharmacy settings in LMICs. The synthesis demonstrates that structured, multifaceted interventions targeting retail dispensers substantially reduce non-prescription antibiotic dispensing (Pooled aOR = 0.164, 95% CI: 0.102–0.265), corresponding to an 83.6% reduction in dispensing odds. The consistency of effect sizes across Southeast Asia and Sub-Saharan Africa (I² = 0.0%) highlights the widespread applicability of pharmacy-directed stewardship models.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Comparison with Prior Systematic Reviews: Previous systematic syntheses, including Afari-Asiedu et al. (2022) and the Cochrane review (2025), concluded that heterogeneity in study designs, settings, and outcome measures prohibited statistical pooling. Our findings demonstrate that when the analytical scope is strictly refined to private retail drug outlets evaluated via objective simulated client methods, intervention effects exhibit remarkable consistency. In private commercial settings, dispensers face common structural drivers — customer demand, fear of financial revenue loss, and weak regulatory enforcement. Multifaceted interventions that address knowledge gaps while providing diagnostic tools (such as point-of-care CRP testing) and peer-group reinforcement effectively mitigate these commercial pressures.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Policy and Practical Implications: These findings indicate that relying solely on top-down legislative prohibitions is insufficient to curb non-prescription antibiotic sales in LMICs, where enforcement infrastructure is often limited. Instead, integrating multifaceted stewardship packages — combining continuing dispenser education, rapid diagnostic aids, customer public awareness campaigns, and peer supervision networks — provides a feasible, high-impact strategy for national AMR containment action plans.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_p("Study Limitations: Several limitations should be considered when interpreting these findings. First, the primary quantitative meta-analysis is based on a small pool of controlled trials (k = 3), reflecting the scarcity of rigorously designed experimental studies in private LMIC pharmacies. Consequently, tests for statistical heterogeneity (Q and I²) have limited statistical power, and the observed I² = 0.0% should be interpreted as the absence of statistically detectable heterogeneity rather than absolute homogeneity. Second, Tumwikirize et al. (2004) — an eligible quasi-experimental trial in Uganda reporting an RR of 0.85 (95% CI: 0.66–1.09) — could not be integrated into the primary log-OR pool due to paywalled full-text access and non-combinable summary statistics. Third, simulated client visits evaluate immediate dispensing behavior during single encounters, which may not capture long-term community consumption trends.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── CONCLUSION ──
add_p("Conclusion", bold=True, size=12, space_before=6, space_after=6)

add_p("Multifaceted interventions integrating dispenser education, rapid diagnostic testing, customer awareness, and peer supervision significantly reduce non-prescription antibiotic dispensing in LMIC community pharmacies (Pooled aOR = 0.164, 95% CI: 0.102–0.265). National health authorities and global AMR stewardship programs should prioritize the adaptation and scaled implementation of community pharmacy-focused intervention models, supported by rigorous multi-country cluster randomized trials to evaluate long-term sustainability.",
      align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)

# ── REFERENCES ──
add_p("References", bold=True, size=12, space_before=6, space_after=6)

references_list = [
    "Murray CJL, Ikuta KS, Sharara F, Swetschinski LR, Robles Aguilar G, Gray A, et al. Global burden of bacterial antimicrobial resistance in 2019: a systematic analysis. Lancet. 2022;399(10325):629–55.",
    "Morgan DJ, Okeke IN, Laxminarayan R, Perencevich EN, Weisenberg S. Non-prescription antimicrobial use worldwide: a systematic review. Lancet Infect Dis. 2011;11(9):692–701.",
    "Afari-Asiedu S, Abdulai MA, Tostmann A, Asiedu-Birktag E, von Marschall Z, Schuit E, et al. Interventions to improve dispensing of antibiotics at the community level in low and middle income countries: a systematic review. J Glob Antimicrob Resist. 2022;29:259–74.",
    "Chalker J, Ratanawijitrasin S, Chuc NTK, Petzold M, Tomson G. Effectiveness of a multi-component intervention on dispensing practices at private pharmacies in Vietnam and Thailand—a randomized controlled trial. Soc Sci Med. 2005;60(1):131–41.",
    "Onwunduba A, Ekwunife O, Onyilogwu E, Onyemelukwe C, Modebe A, Igbokwe D. Impact of point-of-care C-reactive protein testing intervention on non-prescription dispensing of antibiotics for respiratory tract infections in private community pharmacies in Nigeria: a cluster randomized controlled trial. Int J Infect Dis. 2023;127:137–43.",
    "Ferdiana A, Wulandari LPL, Liverani M, Limato R, Essiet I, Mcknight J, et al. The impact of a multi-faceted intervention on non-prescription dispensing of antibiotics by urban community pharmacies in Indonesia: a mixed methods evaluation. BMJ Glob Health. 2024;9(10):e015620.",
    "Tumwikirize WA, Ekwaru PJ, Mohammed K, Ogwal-Okeng JW, Aupont O. Impact of a face-to-face educational intervention on improving the management of acute respiratory infections in private pharmacies and drug shops in Uganda. East Afr Med J. 2004;81(Suppl 1):S33–40.",
    "DerSimonian R, Laird N. Meta-analysis in clinical trials. Control Clin Trials. 1986;7(3):177–88."
]

for idx, ref in enumerate(references_list, 1):
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.left_indent = Inches(0.3)
    p_ref.paragraph_format.first_line_indent = Inches(-0.3)
    p_ref.paragraph_format.space_after = Pt(4)
    p_ref.paragraph_format.line_spacing = 1.15
    run = p_ref.add_run(f"{idx}.  {ref}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9.5)

# Save clean file
output_path = "Manuscript_PINMAS_2026_FINAL_SUBMISSION.docx"
doc.save(output_path)
print(f"SUCCESS: Clean manuscript saved to {output_path}")
